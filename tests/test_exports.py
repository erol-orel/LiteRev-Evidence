"""Exports of the relevant articles (api/exports.py): formatters are pure, the endpoint
serves a file with the right type and name. The relevance query is stubbed."""
import json
import re

import main
from conftest import patch_app  # noqa: E402


def _articles():
    return [
        {"id": 7, "title": "Autochthonous chikungunya in Italy {2024}", "authors": "Rossi Maria; Bianchi Luca",
         "year": 2024, "journal": "Eurosurveillance", "doi": "10.2807/x", "pmid": "39000001",
         "study_design": "Cohort", "similarity_score": 0.8123456, "quality_score": 0.71, "citation_count": 3,
         "screening_status": "included", "abstract": "Line one.\nLine two.", "keywords": "Aedes; Italy",
         "pico_json": {"P": "residents", "I": "vector control", "C": "none", "O": "cases"}},
        {"id": 8, "title": "Second article", "authors": None, "year": None, "journal": None, "doi": None,
         "similarity_score": None, "abstract": None, "pico_json": None},
    ]


def test_rows_and_tabular_formats():
    rows = main.export_rows(_articles())
    assert [r["rank"] for r in rows] == [1, 2]
    assert rows[0]["url"] == "https://doi.org/10.2807/x" and rows[0]["similarity_score"] == 0.8123
    assert rows[0]["pico_population"] == "residents" and rows[1]["pico_population"] == ""
    csv_text = main.to_csv(rows)
    assert csv_text.startswith("﻿rank,id,title,authors,year")
    assert "Autochthonous chikungunya in Italy {2024}" in csv_text and csv_text.count("\r\n") == 3
    data = json.loads(main.to_json(rows, {"scenario": "s"}))
    assert data["meta"]["scenario"] == "s" and data["articles"][1]["title"] == "Second article"
    xlsx = main.to_xlsx(rows, "Test")
    assert xlsx[:2] == b"PK"                                   # a zip, hence a real workbook
    md = main.to_markdown(rows, "Chikungunya")
    assert md.startswith("# Chikungunya") and "1. **Autochthonous" in md


def test_ris_and_bibtex_records():
    rows = main.export_rows(_articles())
    ris = main.to_ris(rows)
    assert ris.count("TY  - JOUR") == 2 and ris.count("ER  - ") == 2
    assert "AU  - Rossi Maria\r\nAU  - Bianchi Luca" in ris
    assert "DO  - 10.2807/x" in ris and "AB  - Line one. Line two." in ris and "KW  - Aedes" in ris
    bib = main.to_bibtex(rows)
    assert bib.startswith("@article{rossi2024_7,")
    assert "title = {Autochthonous chikungunya in Italy \\{2024\\}}" in bib   # braces escaped
    assert "author = {Rossi Maria and Bianchi Luca}" in bib
    assert "@article{anonnd_8," in bib


# ── RIS that Zotero actually reads ───────────────────────────────────────────
# Chaque test ci-dessous correspond à un écart qui faisait qu'un fichier exporté
# s'importait mal, ou pas du tout.

def test_the_journal_lands_in_the_publication_field_not_in_the_abbreviation():
    """THE defect: Zotero maps `JO` to the journal ABBREVIATION, so every imported
    reference arrived with an empty publication, the one field a citation needs."""
    ris = main.to_ris(main.export_rows(_articles()))
    assert "T2  - Eurosurveillance" in ris
    assert "JF  - Eurosurveillance" in ris
    assert "JO  - " not in ris


def test_records_are_separated_by_a_blank_line():
    """`ER  - ` immediately followed by the next `TY` is not valid RIS; EndNote and
    Mendeley stop reading there."""
    ris = main.to_ris(main.export_rows(_articles()))
    assert "ER  - \r\n\r\nTY  - " in ris
    assert "ER  - \r\nTY  - " not in ris
    assert ris.endswith("ER  - \r\n")


def test_a_preprint_is_not_exported_as_a_journal_article():
    rows = main.export_rows([{"id": 9, "title": "A preprint", "authors": "Roe B", "year": 2025,
                              "journal": "medRxiv", "source": "medrxiv"}])
    ris = main.to_ris(rows)
    assert "TY  - UNPB" in ris and "TY  - JOUR" not in ris
    # Un dépôt n'est pas une revue : il ne doit pas remplir le titre de publication.
    assert "T2  - " not in ris and "PB  - medRxiv" in ris


def test_a_guideline_is_exported_as_a_report():
    rows = main.export_rows([{"id": 10, "title": "WHO guideline", "authors": "WHO", "year": 2025,
                              "study_design": "Clinical practice guideline", "source": "pubmed"}])
    assert "TY  - RPRT" in main.to_ris(rows)


def test_no_field_can_span_two_lines():
    """A newline left in a title or an abstract splits the record in two and the import
    stops at that point."""
    rows = main.export_rows([{"id": 11, "title": "A title\nwith a break", "authors": "Roe B",
                              "abstract": "First.\n\nSecond.", "keywords": "one\ntwo"}])
    ris = main.to_ris(rows)
    assert "TI  - A title with a break" in ris
    assert "AB  - First. Second." in ris
    for line in ris.split("\r\n"):
        assert line == "" or re.match(r"^[A-Z][A-Z0-9]  - ", line), line


def test_the_doi_is_bare_so_zotero_recognises_it():
    rows = main.export_rows([{"id": 12, "title": "T", "doi": "https://doi.org/10.1000/abc"},
                             {"id": 13, "title": "U", "doi": "doi: 10.1000/def"}])
    ris = main.to_ris(rows)
    assert "DO  - 10.1000/abc" in ris and "DO  - 10.1000/def" in ris
    assert "DO  - https://" not in ris


def test_the_pmid_follows_pubmeds_own_ris_convention():
    ris = main.to_ris(main.export_rows(_articles()))
    assert "DB  - PubMed\r\nAN  - 39000001" in ris


def test_every_record_opens_with_its_type_and_closes_with_er():
    """Zotero detects a RIS file by its first line and ends each item on `ER`."""
    ris = main.to_ris(main.export_rows(_articles()))
    assert ris.startswith("TY  - ")
    for record in ris.split("\r\n\r\n"):
        lines = [ln for ln in record.split("\r\n") if ln]
        assert lines[0].startswith("TY  - ") and lines[-1] == "ER  - "


def test_an_empty_export_is_an_empty_file_not_a_broken_record():
    assert main.to_ris([]) == ""


def test_bibtex_records():
    rows = main.export_rows(_articles())
    bib = main.to_bibtex(rows)
    assert bib.startswith("@article{rossi2024_7,")
    assert "title = {Autochthonous chikungunya in Italy \\{2024\\}}" in bib   # braces escaped
    assert "author = {Rossi Maria and Bianchi Luca}" in bib
    assert "@article{anonnd_8," in bib


def test_the_endpoint_serves_a_named_file(monkeypatch):
    from fastapi.testclient import TestClient

    patch_app(monkeypatch, "_get_user_scenario_or_404", lambda sid: {"id": sid, "name": "Chikungunya Europe", "query": "q"})
    patch_app(monkeypatch, "_get_above_threshold_articles", lambda sid, threshold=None, **k: _articles())
    patch_app(monkeypatch, "_export_extra_fields", lambda ids: {7: {"id": 7, "pmid": "39000001", "source": "pubmed", "url": None, "keywords": "Aedes"}})
    client = TestClient(main.app)
    r = client.get("/user-scenarios/usr-x/relevant/export?format=ris")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/x-research-info-systems")
    assert r.headers["content-disposition"] == 'attachment; filename="chikungunya-europe_relevant-articles_2.ris"'
    assert r.headers["x-article-count"] == "2" and "TY  - JOUR" in r.text

    r = client.get("/gesica/scenarios/influenza-surveillance/relevant/export?format=xlsx&include_abstract=false")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert client.get("/user-scenarios/usr-x/relevant/export?format=docx").status_code == 400


# ── Exports of SUBSETS: clusters, concepts, an explicit selection ────────────
def _client(monkeypatch):
    from fastapi.testclient import TestClient

    patch_app(monkeypatch, "_get_user_scenario_or_404",
              lambda sid: {"id": sid, "name": "Chikungunya Europe", "query": "q"})
    patch_app(monkeypatch, "_export_extra_fields", lambda ids: {})
    patch_app(monkeypatch, "articles_by_ids", lambda sid, ids: [
        a for a in _articles() if int(a["id"]) in {int(i) for i in ids}])
    return TestClient(main.app)


def _cluster_cache():
    return {
        "lang": "fr", "n_docs": 3000, "n_docs_total": 8200,
        "clusters": [
            {"cluster_id": 0, "cluster_name": "Vector control", "n_docs": 2,
             "top_words": ["aedes", "control"], "summary": "Sur la lutte antivectorielle.",
             "points": [{"id": 7, "x": 0.1, "y": 0.2}, {"id": 8, "x": 0.3, "y": 0.4}]},
            {"cluster_id": -1, "cluster_name": "Bruit", "is_noise": True, "points": []},
        ],
    }


def test_a_cluster_exports_its_articles_and_declares_its_bounds(monkeypatch):
    """A cluster export is complete FOR THAT CLUSTER (the cache keeps every point), but
    clustering itself runs on the most relevant CLUSTER_MAX_DOCS only. The file has to
    say so, or it reads as "every article on this theme"."""
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_load_viz_cache", lambda sid, col, ttl=None: _cluster_cache())

    r = client.get("/user-scenarios/usr-x/clusters/0/export?format=json")
    assert r.status_code == 200, r.text
    assert r.headers["x-export-subset"] == "cluster"
    assert r.headers["x-article-count"] == "2"
    assert "vector-control" in r.headers["content-disposition"]
    meta = json.loads(r.text)["meta"]
    assert meta["subset"] == "cluster" and meta["cluster_id"] == 0
    assert meta["n_docs_clustered"] == 3000 and meta["n_docs_eligible"] == 8200
    # The bound is stated in the file itself, not only in the interface.
    assert "3000" in meta["coverage"] and "8200" in meta["coverage"]
    assert "pas le corpus entier" in meta["coverage"]


def test_an_unknown_cluster_and_a_missing_cache_are_both_explained(monkeypatch):
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_load_viz_cache", lambda sid, col, ttl=None: _cluster_cache())
    r = client.get("/user-scenarios/usr-x/clusters/42/export")
    assert r.status_code == 404 and "42" in r.json()["detail"]

    patch_app(monkeypatch, "_load_viz_cache", lambda sid, col, ttl=None: None)
    r = client.get("/user-scenarios/usr-x/clusters/0/export")
    assert r.status_code == 404 and "Clusters" in r.json()["detail"]


def _concept_row(i, *concepts):
    """A relevant article as `_concept_rows` reads it, with its LLM concepts."""
    return {"id": i, "title": f"Article {i}", "year": 2024, "quality": 0.5, "doi": None,
            "pmid": None, "country": None, "study_design": None, "pico_json": None,
            "metadata_json": None, "keywords": None, "similarity": 0.9,
            "concepts_json": {"v": 2, "concepts": [{"t": t, "en": en, "fr": en}
                                                    for t, en in concepts]}}


def _concept_corpus():
    """Article 7: chikungunya and Aedes albopictus; article 8: chikungunya only."""
    return [_concept_row(7, ("pathogen", "chikungunya virus"), ("vector", "Aedes albopictus")),
            _concept_row(8, ("pathogen", "chikungunya virus"))]


def test_a_concept_export_reads_the_full_index_not_the_drawn_map(monkeypatch):
    """Nodes served to the interface carry at most 40 articles, and the map draws at most
    60 concepts. The export resolves the concept on the full index of the relevant
    articles instead, with every article of the concept."""
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_concept_rows", lambda sid: (_concept_corpus(), 150))

    r = client.get("/user-scenarios/usr-x/concepts/export"
                   "?concepts=pathogen:chikungunya virus&format=json")
    assert r.status_code == 200, r.text
    meta = json.loads(r.text)["meta"]
    assert meta["subset"] == "concepts" and meta["concepts"] == ["chikungunya virus"]
    assert meta["n_with_concepts"] == 2 and meta["n_total"] == 150
    assert "index COMPLET" in meta["coverage"]
    assert r.headers["x-article-count"] == "2"
    # Any spelling of the concept that the index folds onto it finds it too.
    r = client.get("/user-scenarios/usr-x/concepts/export?concepts=vector:aedes Albopictus")
    assert r.status_code == 200 and r.headers["x-article-count"] == "1"


def test_a_concept_selection_unions_or_intersects(monkeypatch):
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_concept_rows", lambda sid: (_concept_corpus(), 150))

    q = "concepts=pathogen:chikungunya virus|vector:Aedes albopictus"
    r_any = client.get(f"/user-scenarios/usr-x/concepts/export?{q}&mode=any&format=json")
    r_all = client.get(f"/user-scenarios/usr-x/concepts/export?{q}&mode=all&format=json")
    assert r_any.headers["x-article-count"] == "2"      # union: 7 and 8
    assert r_all.headers["x-article-count"] == "1"      # intersection: only 7
    assert " OU " in json.loads(r_any.text)["meta"]["coverage"]
    assert " ET " in json.loads(r_all.text)["meta"]["coverage"]


def test_an_unknown_concept_is_named_rather_than_silently_dropped(monkeypatch):
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_concept_rows", lambda sid: (_concept_corpus(), 150))
    r = client.get("/user-scenarios/usr-x/concepts/export?concepts=pathogen:zika virus")
    assert r.status_code == 404 and "zika virus" in r.json()["detail"]
    # A malformed selector says what the shape should be.
    r = client.get("/user-scenarios/usr-x/concepts/export?concepts=justalabel")
    assert r.status_code == 400 and "type:label" in r.json()["detail"]
    # And asking for nothing at all is a 400, not the whole index.
    assert client.get("/user-scenarios/usr-x/concepts/export").status_code == 400


def _wide_concept_corpus():
    """Seventy outcomes, each cited by two articles: more than the 60 nodes the map draws."""
    rows = []
    for k in range(70):
        for j in (0, 1):
            rows.append(_concept_row(1000 + 2 * k + j, ("outcome", f"outcome number {k:02d}")))
    return rows


def test_a_concept_beyond_the_drawn_map_is_exported_rather_than_refused(monkeypatch):
    """THE defect: the export looked the concept up among the map's 60 nodes, so a
    concept of the corpus that the map did not draw answered 404."""
    import main as _main

    rows = _wide_concept_corpus()
    drawn = {n["label"]["en"] for n in _main._build_concept_graph(rows)["nodes"]}
    beyond = next(f"outcome number {k:02d}" for k in range(70)
                  if f"outcome number {k:02d}" not in drawn)
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_concept_rows", lambda sid: (rows, len(rows)))
    patch_app(monkeypatch, "articles_by_ids", lambda sid, ids: [
        {"id": int(i), "title": f"Article {i}"} for i in ids])
    r = client.get(f"/user-scenarios/usr-x/concepts/export?concepts=outcome:{beyond}&format=json")
    assert r.status_code == 200, r.text
    assert r.headers["x-article-count"] == "2"


def test_the_composed_map_export_follows_the_filters_not_the_display_cut(monkeypatch):
    """"Export the map as composed": every concept of the index minus what the user hid,
    removed or set below the minimum, never minus what the screen had no room for."""
    rows = _wide_concept_corpus() + [
        _concept_row(5000, ("pathogen", "dengue virus")),            # a single article
        _concept_row(5001, ("vector", "Aedes aegypti")), _concept_row(5002, ("vector", "Aedes aegypti")),
    ]
    client = _client(monkeypatch)
    patch_app(monkeypatch, "_concept_rows", lambda sid: (rows, len(rows)))
    patch_app(monkeypatch, "articles_by_ids", lambda sid, ids: [
        {"id": int(i), "title": f"Article {i}"} for i in ids])

    everything = client.get("/user-scenarios/usr-x/concepts/export?composed=true&format=json")
    assert everything.status_code == 200, everything.text
    assert everything.headers["x-article-count"] == str(len(rows))      # all 143 articles
    meta = json.loads(everything.text)["meta"]
    assert meta["composed"] is True and meta["n_concepts"] == 72
    assert "concept-map" in everything.headers["content-disposition"]

    # Two articles at least: the dengue article goes. Outcomes hidden: their 140 go.
    # Aedes aegypti removed by hand: nothing is left.
    r = client.get("/user-scenarios/usr-x/concepts/export?composed=true&min_count=2"
                   "&exclude_types=outcome&format=json")
    assert r.headers["x-article-count"] == "2"                         # 5001, 5002
    r = client.get("/user-scenarios/usr-x/concepts/export?composed=true&min_count=2"
                   "&exclude_types=outcome&exclude=vector:Aedes aegypti&format=json")
    assert r.headers["x-article-count"] == "0"
    meta = json.loads(r.text)["meta"]
    assert meta["concept_filter"] == {"exclude_types": ["outcome"], "min_count": 2,
                                      "exclude": ["vector:Aedes aegypti"]}
    assert "sans les types outcome" in meta["coverage"] and "au moins 2" in meta["coverage"]


def test_the_generic_id_export_is_scoped_to_the_scenario(monkeypatch):
    """This is what serves the RAG sources and any hand-picked set. An id that does not
    belong to the scenario must not come out, and the file must say some were dropped."""
    client = _client(monkeypatch)
    r = client.get("/user-scenarios/usr-x/articles/export?ids=7,8,99999&label=RAG answer&format=json")
    assert r.status_code == 200, r.text
    assert r.headers["x-export-subset"] == "selection"
    assert r.headers["x-article-count"] == "2"
    assert r.headers["x-missing-ids"] == "1"
    meta = json.loads(r.text)["meta"]
    assert meta["requested"] == 3 and meta["returned"] == 2 and meta["missing_ids"] == [99999]
    assert "écartés" in meta["coverage"]
    assert "rag-answer" in r.headers["content-disposition"]

    assert client.get("/user-scenarios/usr-x/articles/export?ids=").status_code == 400
    assert client.get("/user-scenarios/usr-x/articles/export?ids=abc").status_code == 400


# ── The REAL path: no stubs on the cache loader or the article query ─────────
# The tests above stub `_load_viz_cache` and `articles_by_ids`, so they proved the
# formatters and the plumbing but NOT that the endpoint talks to the database correctly.
# It did not: `_load_viz_cache` takes the cache KEY ("clustering"), and the endpoint
# passed the column name ("clustering_json"), so every cluster export raised KeyError in
# production while the suite stayed green. These tests go through the database.
SID_REAL = "usr-export-realpath"


def _seed_clustered_scenario(db_conn):
    from conftest import ensure_document_columns
    import json as _json

    ensure_document_columns(db_conn.cursor())
    with db_conn.cursor() as cur:
        cur.execute("DELETE FROM article_scenarios WHERE scenario_id = %s", (SID_REAL,))
        cur.execute("DELETE FROM literature_document WHERE id BETWEEN 9200 AND 9209")
        cur.execute("DELETE FROM scenario_settings WHERE scenario_id = %s", (SID_REAL,))
        cur.execute("INSERT INTO user_scenarios (id, name, query, created_at, updated_at) "
                    "VALUES (%s,'Real path','q',NOW(),NOW()) ON CONFLICT (id) DO NOTHING", (SID_REAL,))
        for i in (9200, 9201, 9202):
            cur.execute("INSERT INTO literature_document (id, title, abstract, year, journal, doi, "
                        "source, project_context) VALUES (%s,%s,'abstract',2025,'Circulation',%s,"
                        "'pubmed','literev')", (i, f"Calcification paper {i}", f"10.1/{i}"))
            cur.execute("INSERT INTO article_scenarios (document_id, scenario_id, similarity_score) "
                        "VALUES (%s,%s,0.80)", (i, SID_REAL))
        payload = {
            "scenario_id": SID_REAL, "lang": "fr", "n_docs": 3, "n_docs_total": 8,
            "clusters": [
                {"cluster_id": 0, "cluster_name": "Calcification", "n_docs": 3,
                 "top_words": ["calcification"], "summary": "Sur la calcification.",
                 "points": [{"id": 9200, "x": 0.1, "y": 0.1}, {"id": 9201, "x": 0.2, "y": 0.2},
                            {"id": 9202, "x": 0.3, "y": 0.3}]},
            ],
        }
        cur.execute("INSERT INTO scenario_settings (scenario_id, clustering_json, clustering_generated_at) "
                    "VALUES (%s, CAST(%s AS jsonb), NOW()) ON CONFLICT (scenario_id) DO UPDATE "
                    "SET clustering_json = CAST(%s AS jsonb), clustering_generated_at = NOW()",
                    (SID_REAL, _json.dumps(payload), _json.dumps(payload)))


def test_a_cluster_export_works_against_a_real_clustering_cache(db_conn):
    """End to end: a real scenario, a real clustering cache, a real article query."""
    from fastapi.testclient import TestClient

    _seed_clustered_scenario(db_conn)
    r = TestClient(main.app).get(f"/user-scenarios/{SID_REAL}/clusters/0/export?format=csv")
    assert r.status_code == 200, r.text
    assert r.headers["x-article-count"] == "3"
    body = r.content.decode("utf-8")
    for i in (9200, 9201, 9202):
        assert f"Calcification paper {i}" in body
    assert r.headers["x-export-subset"] == "cluster"


def test_the_real_cluster_export_states_its_projection_bound(db_conn):
    from fastapi.testclient import TestClient

    _seed_clustered_scenario(db_conn)
    r = TestClient(main.app).get(f"/user-scenarios/{SID_REAL}/clusters/0/export?format=json")
    assert r.status_code == 200, r.text
    meta = json.loads(r.text)["meta"]
    assert meta["n_docs_clustered"] == 3 and meta["n_docs_eligible"] == 8
    assert "3" in meta["coverage"] and "8" in meta["coverage"]


def test_the_real_id_export_is_bounded_by_the_scenario(db_conn):
    """`articles_by_ids` against the database: an id from another scenario is dropped."""
    from fastapi.testclient import TestClient

    _seed_clustered_scenario(db_conn)
    r = TestClient(main.app).get(
        f"/user-scenarios/{SID_REAL}/articles/export?ids=9200,9999999&format=json")
    assert r.status_code == 200, r.text
    assert r.headers["x-article-count"] == "1"
    assert r.headers["x-missing-ids"] == "1"


SID_CORPUS = "usr-export-whole-corpus"


def test_the_whole_corpus_export_has_every_article_and_says_which_are_relevant(db_conn):
    """The search page exported the 200 rows of its preview. Its whole-corpus choice is
    every article the search found: below the threshold and excluded at screening
    included, duplicates not; the file says how many of them the analyses read."""
    from conftest import ensure_document_columns
    from fastapi.testclient import TestClient

    ensure_document_columns(db_conn.cursor())
    with db_conn.cursor() as cur:
        cur.execute("DELETE FROM article_scenarios WHERE scenario_id = %s", (SID_CORPUS,))
        cur.execute("DELETE FROM literature_document WHERE id BETWEEN 9210 AND 9219")
        cur.execute("DELETE FROM scenario_settings WHERE scenario_id = %s", (SID_CORPUS,))
        cur.execute("INSERT INTO user_scenarios (id, name, query, created_at, updated_at) "
                    "VALUES (%s,'Whole corpus','q',NOW(),NOW()) ON CONFLICT (id) DO NOTHING",
                    (SID_CORPUS,))
        cur.execute("INSERT INTO scenario_settings (scenario_id, similarity_threshold) VALUES (%s, 0.5)",
                    (SID_CORPUS,))
        # (id, similarity, screening status, duplicate)
        for i, sim, status, dup in ((9210, 0.80, None, False), (9211, 0.10, None, False),
                                    (9212, 0.90, "excluded", False), (9213, 0.20, "included", False),
                                    (9214, 0.95, None, True)):
            cur.execute("INSERT INTO literature_document (id, title, abstract, year, source, "
                        "project_context, is_duplicate) VALUES (%s,%s,'an abstract long enough to count',2024,'pubmed',"
                        "'literev',%s)", (i, f"Corpus paper {i}", dup))
            cur.execute("INSERT INTO article_scenarios (document_id, scenario_id, similarity_score, "
                        "screening_status) VALUES (%s,%s,%s,%s)", (i, SID_CORPUS, sim, status))

    client = TestClient(main.app)
    r = client.get(f"/user-scenarios/{SID_CORPUS}/corpus/export?format=json")
    assert r.status_code == 200, r.text
    assert r.headers["x-article-count"] == "4" and r.headers["x-export-subset"] == "corpus"
    out = json.loads(r.text)
    assert sorted(a["id"] for a in out["articles"]) == [9210, 9211, 9212, 9213]
    assert out["meta"]["n_relevant"] == 2                     # 9210 above, 9213 included
    assert "4 articles" in out["meta"]["coverage"] and "2 sont pertinents" in out["meta"]["coverage"]
    # The relevant export of the same scenario holds those two only.
    rel = client.get(f"/user-scenarios/{SID_CORPUS}/relevant/export?format=json")
    assert sorted(a["id"] for a in json.loads(rel.text)["articles"]) == [9210, 9213]
    assert client.get(f"/gesica/scenarios/{SID_CORPUS}/corpus/export?format=csv").status_code == 200
    assert client.get(f"/user-scenarios/{SID_CORPUS}/corpus/export?format=doc").status_code == 400


def test_every_export_route_exists_for_built_in_scenarios_too():
    """`scenarioBase` sends built-in scenarios to /gesica/scenarios. The relevant export
    had both routes; the three new ones had only the user-scenario one, so their buttons
    pointed at routes that did not exist."""
    paths = {r.path for r in main.app.routes if hasattr(r, "path")}
    for suffix in ("/clusters/{cluster_id}/export", "/concepts/export", "/articles/export",
                   "/relevant/export", "/corpus/export"):
        assert "/user-scenarios/{scenario_id}" + suffix in paths, suffix
        assert "/gesica/scenarios/{scenario_id}" + suffix in paths, f"gesica {suffix}"


def test_a_doi_already_stored_as_a_url_does_not_become_a_double_link():
    """`https://doi.org/https://doi.org/10.…` was a dead link in every format."""
    rows = main.export_rows([{"id": 14, "title": "T", "doi": "https://doi.org/10.1101/2025.01.01"}])
    assert rows[0]["url"] == "https://doi.org/10.1101/2025.01.01"


# ── What every row carries ───────────────────────────────────────────────────
def _scored():
    return [{"id": 20, "title": "Scored", "authors": "Roe B", "year": 2024,
             "similarity_score": 0.61234, "rerank_score": 0.912345, "language": "eng",
             "keywords": "; ".join(f"kw{i}" for i in range(25))}]


def test_the_rerank_score_that_orders_the_rows_is_a_column():
    """The rows are ordered by rerank first; a file that sorts on a score it does not
    carry cannot be checked against the screen."""
    rows = main.export_rows(_scored())
    assert rows[0]["rerank_score"] == 0.9123 and rows[0]["similarity_score"] == 0.6123
    header = main.to_csv(rows).lstrip("﻿").split("\r\n")[0].split(",")
    assert header.index("rerank_score") == header.index("similarity_score") + 1
    assert "0.9123" in main.to_csv(rows)


def test_ris_writes_the_language_and_every_keyword():
    ris = main.to_ris(main.export_rows(_scored()))
    assert "LA  - eng" in ris
    assert ris.count("KW  - ") == 25 and "KW  - kw24" in ris      # it stopped at twenty


def test_the_cosine_is_called_a_similarity_not_the_relevance():
    rows = main.export_rows(_scored())
    ris, bib, md = main.to_ris(rows), main.to_bibtex(rows), main.to_markdown(rows)
    for text_ in (ris, bib, md):
        assert "cosine similarity 0.6123" in text_ and "rerank score 0.9123" in text_
        assert "relevance 0.6" not in text_


def test_an_xlsx_survives_control_characters_and_a_slash_in_its_title():
    """openpyxl rejects \\x00-\\x1f control characters (bar tab and newlines) in a cell,
    and []:*?/\\ in a sheet title: ONE such title or abstract failed the whole export."""
    rows = main.export_rows([{"id": 21, "title": "Broken\x0btitle\x1f", "authors": "Roe B",
                              "abstract": "Line\x00one\x08.\nLine two."}])
    xlsx = main.to_xlsx(rows, "H5N1 / A(H5N1): field [notes]?", ["Scenario: A/H5N1\x07"])
    assert xlsx[:2] == b"PK"
    import io

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=head.index("title") + 1).value == "Brokentitle"
    assert ws.cell(row=2, column=head.index("abstract") + 1).value == "Lineone.\nLine two."


def test_a_long_id_list_goes_in_the_body(monkeypatch):
    """Behind nginx a GET query string of a few thousand ids fails with a 414. The POST
    variant takes the same list in its body and gives the same file."""
    client = _client(monkeypatch)
    ids = [7, 8] + list(range(100000, 101500))
    r = client.post("/user-scenarios/usr-x/articles/export",
                    json={"ids": ids, "format": "json", "label": "Filtered map"})
    assert r.status_code == 200, r.text
    assert r.headers["x-article-count"] == "2" and r.headers["x-missing-ids"] == "1500"
    meta = json.loads(r.text)["meta"]
    assert meta["requested"] == 1502 and meta["returned"] == 2
    assert "filtered-map" in r.headers["content-disposition"]
    # Same answers as the GET for the same small list, and the same refusals.
    small = client.post("/user-scenarios/usr-x/articles/export", json={"ids": [7, 8], "format": "csv"})
    assert small.headers["x-article-count"] == "2"
    assert client.post("/user-scenarios/usr-x/articles/export", json={"ids": []}).status_code == 400
    assert client.post("/user-scenarios/usr-x/articles/export", json={"ids": ["x"]}).status_code == 400
    assert client.post("/user-scenarios/usr-x/articles/export",
                       json={"ids": list(range(20001))}).status_code == 400
    assert client.post("/gesica/scenarios/influenza-surveillance/articles/export",
                       json={"ids": [7]}).status_code == 200
