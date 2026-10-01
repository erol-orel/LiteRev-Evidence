#!/usr/bin/env python3
"""Build and score a gold standard for the epidemiological parameter extraction.

The pipeline that turns literature into SEIR parameters has two stages, and a validation
that measures only one of them measures the easy one:

  1. THE SCREEN. A regular expression over title and abstract (`_PARAM_PHRASES`) decides
     which articles the extraction will even read. Articles it rejects are unreachable:
     no LLM, however good, recovers them. Its recall is therefore a ceiling on the whole
     pipeline, and it can only be estimated on articles the screen said NO to.
  2. THE EXTRACTION. An LLM reads the accepted articles and returns values. Its precision
     and recall are conditional on the screen having let the article through.

So the sample is STRATIFIED and includes a stratum the screen rejected. Measuring only
stratum A would report the extraction's accuracy as if it were the pipeline's.

Three steps, each resumable, each writing files you can archive next to a paper:

  sample   draw and FREEZE the sample (seed, regex fingerprint, article ids, strata)
  sheets   write one blind annotation workbook per annotator, no model output in it
  score    read the filled workbooks plus the extraction output, and report

Blinding: the sheets never show the stratum, the matched phrase, or anything the
extraction produced. An annotator shown the phrase that triggered the screen cannot give
an independent reading of whether the article reports a value.

  python3 scripts/gold_standard.py sample --scenario usr-xxxx --out gold/flu --n 150 --n-negative 50
  python3 scripts/gold_standard.py sheets --out gold/flu --annotators EO,CB
  # annotate gold/flu/annotation_EO.xlsx and annotation_CB.xlsx by hand, then:
  python3 scripts/gold_standard.py score  --out gold/flu --extraction gold/flu/extraction.json
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import random
import re
import sys
from datetime import datetime, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

#: The annotator writes a number, or one of these, in each parameter cell.
NOT_REPORTED = ""          # the abstract says nothing about this parameter
MENTION_ONLY = "M"         # named, but no value given ("we discuss the reproduction number")

#: A value counts as agreeing with the gold one within this relative tolerance. Extraction
#: and annotator can legitimately differ on rounding or on which of several reported
#: estimates is "the" value; they cannot legitimately differ by a factor.
VALUE_TOLERANCE = 0.10


def _params() -> list[str]:
    from api.variables import _PARAM_PHRASES
    return list(_PARAM_PHRASES)


def _regex_fingerprint() -> str:
    """Identifies the screen this sample was drawn against. Changing `_PARAM_PHRASES`
    changes which articles are positives, so a sample drawn against another version is
    not comparable and the score step refuses to pretend otherwise."""
    from api.variables import _param_regex
    return hashlib.sha256(_param_regex().encode("utf-8")).hexdigest()[:16]


# ─── sample ──────────────────────────────────────────────────────────────────
def _screen_negative_articles(scenario_id: str, threshold: float | None) -> list[dict]:
    """Relevant articles the screen REJECTED. The pipeline can never see these, so they
    are where its ceiling is measured."""
    from sqlalchemy import text

    from api.core import engine
    from api.scenario_store import _get_scenario_threshold, relevant_gate_sql
    from api.variables import _param_regex

    thr = _get_scenario_threshold(scenario_id) if threshold is None else float(threshold)
    gate = relevant_gate_sql(doc="d", link="ars", thr=":thr")
    with engine.connect() as conn:
        rows = conn.execute(text(f"""
            SELECT d.id, d.title, d.abstract, d.year, d.doi, d.study_design, d.source
            FROM literature_document d
            JOIN article_scenarios ars ON ars.document_id = d.id
            WHERE ars.scenario_id = :sid
              AND d.abstract IS NOT NULL
              AND {gate}
              AND NOT ((d.title || ' ' || d.abstract) ~* :rx)
            ORDER BY d.id
        """), {"sid": scenario_id, "thr": thr,
               "rx": _param_regex(boundary=r"\y")}).mappings().all()
    return [dict(r) for r in rows]


def cmd_sample(args: argparse.Namespace) -> int:
    from api.variables import _parameter_candidate_articles

    rng = random.Random(args.seed)
    pos = _parameter_candidate_articles(args.scenario, args.threshold, 0)
    neg = _screen_negative_articles(args.scenario, args.threshold)
    print(f"relevant corpus: {len(pos)} screened IN, {len(neg)} screened OUT")
    if not pos:
        print("nothing to sample: the screen accepted no article in this scenario")
        return 1

    # Stratum A, by parameter: a simple random draw would under-represent the rare
    # parameters (immunity duration, infectious period), and those are exactly the ones
    # whose extraction nobody has checked.
    by_param: dict[str, list[dict]] = {p: [] for p in _params()}
    for a in pos:
        for p in a.get("params_mentioned") or []:
            by_param.setdefault(p, []).append(a)
    per = max(1, args.n // max(1, len([p for p, v in by_param.items() if v])))
    picked: dict[int, dict] = {}
    for p, arts in by_param.items():
        if not arts:
            print(f"  ! no article mentions {p}: it cannot be validated from this corpus")
            continue
        for a in rng.sample(arts, min(per, len(arts))):
            picked.setdefault(int(a["id"]), a)
    # Top the stratum up to n with a plain random draw over the rest.
    rest = [a for a in pos if int(a["id"]) not in picked]
    rng.shuffle(rest)
    for a in rest[:max(0, args.n - len(picked))]:
        picked[int(a["id"])] = a

    neg_pick = rng.sample(neg, min(args.n_negative, len(neg))) if neg else []
    if len(neg_pick) < args.n_negative:
        print(f"  ! only {len(neg_pick)} screened-out articles available, asked {args.n_negative}")

    rows = ([{"id": int(a["id"]), "stratum": "screened_in",
              "params_mentioned": a.get("params_mentioned") or [],
              "title": a.get("title"), "abstract": a.get("abstract"), "year": a.get("year"),
              "doi": a.get("doi"), "study_design": a.get("study_design"),
              "source": a.get("source")} for a in picked.values()]
            + [{"id": int(a["id"]), "stratum": "screened_out", "params_mentioned": [],
                "title": a.get("title"), "abstract": a.get("abstract"), "year": a.get("year"),
                "doi": a.get("doi"), "study_design": a.get("study_design"),
                "source": a.get("source")} for a in neg_pick])
    # Shuffled ONCE here and kept in this order everywhere after: the annotator must not
    # be able to infer the stratum from the position of a row.
    rng.shuffle(rows)

    out = {
        "scenario_id": args.scenario,
        "threshold": args.threshold,
        "seed": args.seed,
        "drawn_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "regex_fingerprint": _regex_fingerprint(),
        "parameters": _params(),
        "population": {"screened_in": len(pos), "screened_out": len(neg)},
        "sampled": {"screened_in": sum(1 for r in rows if r["stratum"] == "screened_in"),
                    "screened_out": sum(1 for r in rows if r["stratum"] == "screened_out")},
        "articles": rows,
    }
    os.makedirs(args.out, exist_ok=True)
    path = os.path.join(args.out, "sample.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=2)
    print(f"wrote {path}: {out['sampled']['screened_in']} screened in, "
          f"{out['sampled']['screened_out']} screened out, seed {args.seed}, "
          f"regex {out['regex_fingerprint']}")
    print("\nSample it ONCE. Re-running with a different seed after seeing results is how"
          "\na validation stops being one.")
    return 0


# ─── sheets ──────────────────────────────────────────────────────────────────
_INSTRUCTIONS = [
    ("What to do",
     "For each article, read the title and abstract and record what the article ITSELF "
     "reports for each of the six parameters. One row per article."),
    ("Leave the cell EMPTY",
     "if the abstract says nothing about that parameter."),
    (f"Write {MENTION_ONLY!r}",
     "if the parameter is named but no value is given (for example 'we discuss the "
     "reproduction number', or 'the serial interval remains uncertain')."),
    ("Write the NUMBER",
     "if the abstract states a value. Use the central value when a range or interval is "
     "given. Units: r0 is a plain ratio; serial interval, incubation period, infectious "
     "period and immunity duration are in DAYS; cfr is a PROPORTION, so a case fatality "
     "of 1.5 percent is 0.015."),
    ("Only this study",
     "Record a value only if the article reports it for the disease it studies. A value "
     "quoted from another paper, or for a different disease, is NOT a value for this "
     "article: leave the cell empty or write M."),
    ("Do not look anything up",
     "The question is what this abstract says, not what is true."),
    ("Disease",
     "Name the disease the article is about, if you can tell."),
    ("Notes",
     "Anything ambiguous. Disagreements get read, so a sentence here is worth having."),
    ("Do not discuss",
     "Annotate independently. The agreement between annotators is itself a measurement, "
     "and it is destroyed by comparing notes before scoring."),
]


def cmd_sheets(args: argparse.Namespace) -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with open(os.path.join(args.out, "sample.json"), encoding="utf-8") as fh:
        sample = json.load(fh)
    params = sample["parameters"]

    for who in [a.strip() for a in args.annotators.split(",") if a.strip()]:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Instructions"
        ws["A1"] = f"Gold standard annotation - {who}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"Scenario {sample['scenario_id']}, {len(sample['articles'])} articles, "
                    f"drawn {sample['drawn_at']}")
        for i, (head, body) in enumerate(_INSTRUCTIONS, start=4):
            ws.cell(row=i, column=1, value=head).font = Font(bold=True)
            c = ws.cell(row=i, column=2, value=body)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 100
        for i in range(4, 4 + len(_INSTRUCTIONS)):
            ws.row_dimensions[i].height = 46

        sh = wb.create_sheet("Annotation")
        head = ["id", "title", "year", "abstract"] + params + ["disease", "notes"]
        sh.append(head)
        for j in range(1, len(head) + 1):
            c = sh.cell(row=1, column=j)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", start_color="DDEBF7")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "E2"
        # The order is the sample's, which was shuffled once: nothing here reveals the
        # stratum, the matched phrase, or anything the extraction produced.
        for a in sample["articles"]:
            sh.append([a["id"], a.get("title") or "", a.get("year") or "",
                       a.get("abstract") or ""] + [""] * len(params) + ["", ""])
        widths = [9, 46, 7, 110] + [13] * len(params) + [20, 30]
        for j, w in enumerate(widths, start=1):
            sh.column_dimensions[get_column_letter(j)].width = w
        for r in range(2, len(sample["articles"]) + 2):
            sh.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
            sh.row_dimensions[r].height = 120

        path = os.path.join(args.out, f"annotation_{who}.xlsx")
        wb.save(path)
        print(f"wrote {path}")
    print("\nSend each annotator only their own file.")
    return 0


# ─── score ───────────────────────────────────────────────────────────────────
def _read_annotation(path: str, params: list[str]) -> dict[int, dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    sh = wb["Annotation"]
    head = [str(c.value or "").strip() for c in sh[1]]
    idx = {name: head.index(name) for name in (["id"] + params + ["disease", "notes"])
           if name in head}
    out: dict[int, dict] = {}
    for row in sh.iter_rows(min_row=2, values_only=True):
        if not row or row[idx["id"]] is None:
            continue
        aid = int(row[idx["id"]])
        rec: dict = {"disease": row[idx["disease"]] if "disease" in idx else None,
                     "notes": row[idx["notes"]] if "notes" in idx else None, "values": {}}
        for p in params:
            if p not in idx:
                continue
            raw = row[idx[p]]
            if raw is None or str(raw).strip() == "":
                rec["values"][p] = ("none", None)
            elif str(raw).strip().upper() == MENTION_ONLY:
                rec["values"][p] = ("mention", None)
            else:
                try:
                    rec["values"][p] = ("value", float(str(raw).strip().replace(",", ".")))
                except ValueError:
                    rec["values"][p] = ("unparsed", str(raw).strip())
        out[aid] = rec
    return out


def _extraction_by_article(extraction: dict, params: list[str]) -> dict[int, dict]:
    """`extract_epidemic_observations` output, inverted to article -> {param: value}."""
    out: dict[int, dict] = {}
    for name, blob in (extraction.get("params") or {}).items():
        if name not in params:
            continue
        for obs in blob.get("observations") or []:
            try:
                aid = int(obs["article_id"])
            except (KeyError, TypeError, ValueError):
                continue
            out.setdefault(aid, {})[name] = obs.get("value")
    return out


def score_parameter(gold: dict, pred: dict, ids: list[int], param: str,
                    tolerance: float = VALUE_TOLERANCE) -> dict:
    """One parameter's confusion matrix and value accuracy. PURE, tested offline.

    `gold[id]["values"][param]` is (kind, value) with kind in none/mention/value/disputed;
    `pred[id][param]` is the extracted number or absent.

    A gold `mention` counts as a NEGATIVE: the article named the parameter without giving
    a value, so an extraction that returns a number there has invented one. That is the
    single most important cell in the table, and collapsing mention into "reported" would
    hide it. A `disputed` cell is excluded entirely: the annotators did not agree, so it
    is not evidence either way."""
    tp = fp = fn = 0
    errors: list[float] = []
    for i in ids:
        kind, gval = (gold.get(i, {}).get("values") or {}).get(param, ("none", None))
        pval = (pred.get(i) or {}).get(param)
        if kind == "disputed":
            continue
        if kind == "value" and pval is not None:
            tp += 1
            if gval:
                errors.append(abs(float(pval) - float(gval)) / abs(float(gval)))
        elif kind == "value":
            fn += 1
        elif pval is not None:
            fp += 1
    prec = tp / (tp + fp) if tp + fp else None
    rec = tp / (tp + fn) if tp + fn else None
    f1 = (2 * prec * rec / (prec + rec)) if prec and rec and (prec + rec) else None
    within = (sum(1 for e in errors if e <= tolerance) / len(errors)) if errors else None
    return {
        "tp": tp, "fp": fp, "fn": fn,
        "precision": None if prec is None else round(prec, 3),
        "recall": None if rec is None else round(rec, 3),
        "f1": None if f1 is None else round(f1, 3),
        "n_values_compared": len(errors),
        "within_tolerance": None if within is None else round(within, 3),
        "median_relative_error": (round(sorted(errors)[len(errors) // 2], 3) if errors else None),
    }


def _kappa(a: list[bool], b: list[bool]) -> float | None:
    """Cohen's kappa on two binary readings. None when it is undefined, which happens
    when both annotators used a single category throughout."""
    n = len(a)
    if n == 0:
        return None
    agree = sum(1 for x, y in zip(a, b) if x == y) / n
    pa, pb = sum(a) / n, sum(b) / n
    chance = pa * pb + (1 - pa) * (1 - pb)
    if abs(1 - chance) < 1e-12:
        return None
    return (agree - chance) / (1 - chance)


def cmd_score(args: argparse.Namespace) -> int:
    with open(os.path.join(args.out, "sample.json"), encoding="utf-8") as fh:
        sample = json.load(fh)
    params = sample["parameters"]
    if sample.get("regex_fingerprint") != _regex_fingerprint():
        print(f"REFUSING: the sample was drawn against screen {sample.get('regex_fingerprint')}, "
              f"the code now has {_regex_fingerprint()}.\nThe phrase list changed, so which "
              f"articles are positives changed with it. Re-sample, or check out the code the "
              f"sample was drawn against.")
        return 2

    sheets = sorted(f for f in os.listdir(args.out)
                    if f.startswith("annotation_") and f.endswith(".xlsx"))
    if not sheets:
        print(f"no annotation_*.xlsx in {args.out}")
        return 1
    annotations = {f[len("annotation_"):-len(".xlsx")]: _read_annotation(os.path.join(args.out, f), params)
                   for f in sheets}
    print(f"annotators: {', '.join(annotations)}")

    strata = {int(a["id"]): a["stratum"] for a in sample["articles"]}
    ids = [int(a["id"]) for a in sample["articles"]]

    # ── agreement between annotators, before anything is merged ──────────────
    names = list(annotations)
    report: dict = {"scenario_id": sample["scenario_id"], "sample": sample["sampled"],
                    "annotators": names, "agreement": {}, "screen": {}, "extraction": {},
                    "scored_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
    if len(names) >= 2:
        A, B = annotations[names[0]], annotations[names[1]]
        common = [i for i in ids if i in A and i in B]
        print(f"\n── Agreement between {names[0]} and {names[1]} on {len(common)} articles")
        for p in params:
            ra = [A[i]["values"].get(p, ("none", None))[0] == "value" for i in common]
            rb = [B[i]["values"].get(p, ("none", None))[0] == "value" for i in common]
            k = _kappa(ra, rb)
            agree = sum(1 for x, y in zip(ra, rb) if x == y) / len(common) if common else 0
            report["agreement"][p] = {"n": len(common), "percent_agreement": round(agree, 3),
                                      "kappa": None if k is None else round(k, 3),
                                      "n_reported": {names[0]: sum(ra), names[1]: sum(rb)}}
            print(f"   {p:26} agreement {agree:5.1%}  kappa "
                  f"{'n/a' if k is None else f'{k:5.2f}'}  "
                  f"({sum(ra)} vs {sum(rb)} reported)")
        if args.adjudicated and os.path.exists(os.path.join(args.out, args.adjudicated)):
            gold = _read_annotation(os.path.join(args.out, args.adjudicated), params)
            print(f"\ngold = {args.adjudicated} (adjudicated)")
        else:
            # Until the disagreements are adjudicated, the defensible gold is the
            # agreement: a pair they differ on is not evidence of anything.
            gold = {}
            dropped = 0
            for i in common:
                vals = {}
                for p in params:
                    ka, va = A[i]["values"].get(p, ("none", None))
                    kb, vb = B[i]["values"].get(p, ("none", None))
                    same = ka == kb and (ka != "value" or _close(va, vb))
                    if same:
                        vals[p] = (ka, va)
                    else:
                        vals[p] = ("disputed", None)
                        dropped += 1
                gold[i] = {"values": vals}
            print(f"gold = the two annotators' agreement; {dropped} parameter cells "
                  f"disputed and excluded. Adjudicate them and pass --adjudicated to "
                  f"score on the full sample.")
    else:
        gold = annotations[names[0]]
        print("\nONE annotator only: no agreement can be measured, and a single reading "
              "is not a gold standard. Add a second.")

    # ── the screen's ceiling, measured on what it rejected ───────────────────
    out_ids = [i for i in ids if strata.get(i) == "screened_out" and i in gold]
    missed = [i for i in out_ids
              if any(gold[i]["values"].get(p, ("none", None))[0] == "value" for p in params)]
    in_ids = [i for i in ids if strata.get(i) == "screened_in" and i in gold]
    in_with_value = [i for i in in_ids
                     if any(gold[i]["values"].get(p, ("none", None))[0] == "value" for p in params)]
    pop = sample["population"]
    print(f"\n── The screen (regex over title and abstract)")
    print(f"   accepted {pop['screened_in']} of {pop['screened_in'] + pop['screened_out']} "
          f"relevant articles")
    if out_ids:
        rate = len(missed) / len(out_ids)
        est = rate * pop["screened_out"]
        print(f"   of {len(out_ids)} it REJECTED, {len(missed)} do report a value "
              f"({rate:.1%})")
        print(f"   so it misses an estimated {est:.0f} articles across the corpus, and its "
              f"recall is about {pop['screened_in'] * (len(in_with_value) / max(1, len(in_ids))) / max(1e-9, pop['screened_in'] * (len(in_with_value) / max(1, len(in_ids))) + est):.1%}")
        report["screen"] = {"n_rejected_checked": len(out_ids), "n_rejected_with_value": len(missed),
                            "miss_rate": round(rate, 4), "estimated_missed_in_corpus": round(est)}
    else:
        print("   no screened-out article annotated: the ceiling is UNMEASURED, and the "
              "extraction numbers below describe only the articles the screen let through")
    if in_ids:
        prec = len(in_with_value) / len(in_ids)
        print(f"   of {len(in_ids)} it ACCEPTED, {len(in_with_value)} do report a value "
              f"({prec:.1%}); the rest are mentions without a value")
        report["screen"]["accepted_precision"] = round(prec, 4)

    # ── the extraction, on what the screen accepted ──────────────────────────
    if not args.extraction:
        print("\nno --extraction given: run the extraction on this scenario, save its JSON, "
              "and pass it to score the second stage")
        _write_report(args.out, report)
        return 0
    with open(args.extraction, encoding="utf-8") as fh:
        pred = _extraction_by_article(json.load(fh), params)
    print(f"\n── The extraction, on the {len(in_ids)} screened-in articles annotated")
    print(f"   {'parameter':26} {'TP':>4} {'FP':>4} {'FN':>4} {'prec':>7} {'recall':>7} {'F1':>7}")
    for p in params:
        r = score_parameter(gold, pred, in_ids, p)
        report["extraction"][p] = r
        _f = lambda x: "    n/a" if x is None else f"{x:7.3f}"            # noqa: E731
        print(f"   {p:26} {r['tp']:>4} {r['fp']:>4} {r['fn']:>4} "
              f"{_f(r['precision'])} {_f(r['recall'])} {_f(r['f1'])}"
              + (f"   values within {VALUE_TOLERANCE:.0%}: {r['within_tolerance']:.0%} "
                 f"(n={r['n_values_compared']})" if r["within_tolerance"] is not None else ""))

    _write_report(args.out, report)
    return 0


def _close(a, b) -> bool:
    if a is None or b is None:
        return a is b
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return a == b
    return abs(a - b) <= VALUE_TOLERANCE * max(abs(a), abs(b), 1e-9)


def _write_report(out_dir: str, report: dict) -> None:
    path = os.path.join(out_dir, "report.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)
    print(f"\nwrote {path}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("sample", help="draw and freeze the sample")
    s.add_argument("--scenario", required=True)
    s.add_argument("--out", required=True, help="directory to write the sample into")
    s.add_argument("--n", type=int, default=150, help="articles the screen ACCEPTED (default 150)")
    s.add_argument("--n-negative", type=int, default=50,
                   help="articles the screen REJECTED, to measure its ceiling (default 50)")
    s.add_argument("--threshold", type=float, default=None)
    s.add_argument("--seed", type=int, default=20260101)
    s.set_defaults(func=cmd_sample)

    h = sub.add_parser("sheets", help="write the blind annotation workbooks")
    h.add_argument("--out", required=True)
    h.add_argument("--annotators", required=True, help="comma-separated initials, e.g. EO,CB")
    h.set_defaults(func=cmd_sheets)

    c = sub.add_parser("score", help="score the filled workbooks")
    c.add_argument("--out", required=True)
    c.add_argument("--extraction", default=None,
                   help="JSON from extract_epidemic_observations for this scenario")
    c.add_argument("--adjudicated", default=None,
                   help="filename of an adjudicated workbook to use as gold instead of "
                        "the annotators' agreement")
    c.set_defaults(func=cmd_score)

    args = ap.parse_args()
    os.environ.setdefault("OPENAI_API_KEY", "")
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
