"""Exports of a scenario's relevant articles: CSV, Excel, RIS, BibTeX, JSON, Markdown.

The relevant articles are the same set every tab uses (`_get_above_threshold_articles`:
included by a reviewer, or scored above the scenario threshold, never the excluded).
The formatters are pure functions over plain dicts, tested offline; the endpoint adds
the bibliographic fields the relevance query does not carry (PMID, source, URL, keywords).
"""
from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

from fastapi import HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import text

from .core import app, engine, logger
from .scenario_store import _get_user_scenario_or_404
from .relevance import _get_above_threshold_articles

EXPORT_FORMATS = ("csv", "xlsx", "ris", "bibtex", "json", "md")
_CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "ris": "application/x-research-info-systems; charset=utf-8",
    "bibtex": "application/x-bibtex; charset=utf-8",
    "json": "application/json; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
}
_EXTENSIONS = {"csv": "csv", "xlsx": "xlsx", "ris": "ris", "bibtex": "bib", "json": "json", "md": "md"}

# One row per article, the same columns in every tabular format.
EXPORT_COLUMNS = (
    "rank", "id", "title", "authors", "year", "journal", "doi", "pmid", "url", "source",
    "study_design", "similarity_score", "quality_score", "citation_count", "screening_status",
    "keywords", "pico_population", "pico_intervention", "pico_comparator", "pico_outcome", "abstract",
)


def _pico(a: dict, key: str) -> str:
    pj = a.get("pico_json")
    if isinstance(pj, str):
        try:
            pj = json.loads(pj)
        except Exception:
            pj = None
    return str((pj or {}).get(key) or "") if isinstance(pj, dict) else ""


def _article_url(a: dict) -> str:
    if a.get("url"):
        return str(a["url"])
    if a.get("doi"):
        return f"https://doi.org/{a['doi']}"
    if a.get("pmid"):
        return f"https://pubmed.ncbi.nlm.nih.gov/{a['pmid']}/"
    return ""


def export_rows(articles: list[dict], include_abstract: bool = True) -> list[dict]:
    """Plain export rows (strings and numbers only), in relevance order. Pure."""
    rows = []
    for i, a in enumerate(articles, 1):
        rows.append({
            "rank": i,
            "id": a.get("id"),
            "title": (a.get("title") or "").strip(),
            "authors": (a.get("authors") or "").strip(),
            "year": a.get("year"),
            "journal": (a.get("journal") or "").strip(),
            "doi": (a.get("doi") or "").strip(),
            "pmid": str(a.get("pmid") or "").strip(),
            "url": _article_url(a),
            "source": (a.get("source") or "").strip(),
            "study_design": (a.get("study_design") or _pico(a, "study_design") or "").strip(),
            "similarity_score": (round(float(a["similarity_score"]), 4)
                                 if a.get("similarity_score") is not None else None),
            "quality_score": (round(float(a["quality_score"]), 3)
                              if a.get("quality_score") is not None else None),
            "citation_count": a.get("citation_count"),
            "screening_status": a.get("screening_status") or "",
            "keywords": (a.get("keywords") or "").strip() if isinstance(a.get("keywords"), str)
                        else "; ".join(str(k) for k in (a.get("keywords") or [])),
            "pico_population": _pico(a, "P"),
            "pico_intervention": _pico(a, "I"),
            "pico_comparator": _pico(a, "C"),
            "pico_outcome": _pico(a, "O"),
            "abstract": (a.get("abstract") or "").strip() if include_abstract else "",
        })
    return rows


def to_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(EXPORT_COLUMNS), extrasaction="ignore", lineterminator="\r\n")
    w.writeheader()
    for r in rows:
        w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in EXPORT_COLUMNS})
    return "﻿" + buf.getvalue()          # BOM: Excel opens UTF-8 accents correctly


def to_json(rows: list[dict], meta: dict | None = None) -> str:
    return json.dumps({"meta": meta or {}, "articles": rows}, ensure_ascii=False, indent=2, default=str)


def _authors_list(authors: str) -> list[str]:
    return [p.strip() for p in re.split(r"[;\n]|,\s(?=[A-Z][a-z]+\s[A-Z])", authors or "") if p.strip()]


def to_ris(rows: list[dict]) -> str:
    """RIS for Zotero, EndNote, Mendeley: one record per article, JOUR type."""
    out = []
    for r in rows:
        lines = ["TY  - JOUR", f"TI  - {r['title']}"]
        for au in _authors_list(r["authors"]):
            lines.append(f"AU  - {au}")
        if r.get("year"):
            lines.append(f"PY  - {r['year']}")
        if r.get("journal"):
            lines.append(f"JO  - {r['journal']}")
        if r.get("doi"):
            lines.append(f"DO  - {r['doi']}")
        if r.get("pmid"):
            lines.append(f"AN  - {r['pmid']}")
        if r.get("url"):
            lines.append(f"UR  - {r['url']}")
        for kw in [k.strip() for k in (r.get("keywords") or "").split(";") if k.strip()][:20]:
            lines.append(f"KW  - {kw}")
        if r.get("abstract"):
            _abstract_one_line = re.sub(r"\s+", " ", r["abstract"])
            lines.append(f"AB  - {_abstract_one_line}")
        notes = []
        if r.get("similarity_score") is not None:
            notes.append(f"relevance {r['similarity_score']}")
        if r.get("study_design"):
            notes.append(f"design {r['study_design']}")
        if notes:
            lines.append(f"N1  - LiteRev: {', '.join(notes)}")
        lines.append("ER  - ")
        out.append("\r\n".join(lines))
    return "\r\n".join(out) + ("\r\n" if out else "")


def _bib_escape(s: str) -> str:
    return (s or "").replace("\\", "\\textbackslash{}").replace("{", "\\{").replace("}", "\\}")


def _bib_key(r: dict) -> str:
    first = (_authors_list(r.get("authors") or "") or ["anon"])[0].split()[0]
    first = re.sub(r"[^A-Za-z]", "", first) or "anon"
    return f"{first.lower()}{r.get('year') or 'nd'}_{r.get('id')}"


def to_bibtex(rows: list[dict]) -> str:
    out = []
    for r in rows:
        fields = [("title", r["title"]), ("author", " and ".join(_authors_list(r["authors"]))),
                  ("year", str(r.get("year") or "")), ("journal", r.get("journal") or ""),
                  ("doi", r.get("doi") or ""), ("url", r.get("url") or ""),
                  ("keywords", r.get("keywords") or ""), ("abstract", r.get("abstract") or ""),
                  ("note", f"LiteRev relevance {r['similarity_score']}" if r.get("similarity_score") is not None else "")]
        body = ",\n".join(f"  {k} = {{{_bib_escape(v)}}}" for k, v in fields if v)
        out.append(f"@article{{{_bib_key(r)},\n{body}\n}}")
    return "\n\n".join(out) + ("\n" if out else "")


def to_markdown(rows: list[dict], title: str = "") -> str:
    lines = [f"# {title}".rstrip(), ""] if title else []
    lines.append(f"{len(rows)} relevant articles, most relevant first.")
    lines.append("")
    for r in rows:
        head = f"{r['rank']}. **{r['title']}**"
        meta = ", ".join(x for x in (r.get("authors"), str(r.get("year") or ""), r.get("journal")) if x)
        link = f" [{r['doi'] or 'link'}]({r['url']})" if r.get("url") else ""
        lines.append(f"{head}{link}")
        if meta:
            lines.append(f"   {meta}")
        extras = []
        if r.get("similarity_score") is not None:
            extras.append(f"relevance {r['similarity_score']}")
        if r.get("study_design"):
            extras.append(r["study_design"])
        if extras:
            lines.append(f"   _{' · '.join(extras)}_")
        lines.append("")
    return "\n".join(lines)


def to_xlsx(rows: list[dict], sheet_title: str = "Relevant articles") -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Articles"
    ws.append(list(EXPORT_COLUMNS))
    for r in rows:
        ws.append([r.get(k) for k in EXPORT_COLUMNS])
    widths = {"title": 60, "authors": 40, "journal": 28, "abstract": 80, "url": 36, "doi": 24,
              "pico_population": 30, "pico_intervention": 30, "pico_comparator": 24, "pico_outcome": 30}
    for i, col in enumerate(EXPORT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_export(fmt: str, rows: list[dict], title: str, meta: dict | None = None) -> bytes:
    """Bytes of the export in `fmt`. Pure (openpyxl for xlsx)."""
    if fmt == "csv":
        return to_csv(rows).encode("utf-8")
    if fmt == "json":
        return to_json(rows, meta).encode("utf-8")
    if fmt == "ris":
        return to_ris(rows).encode("utf-8")
    if fmt == "bibtex":
        return to_bibtex(rows).encode("utf-8")
    if fmt == "md":
        return to_markdown(rows, title).encode("utf-8")
    if fmt == "xlsx":
        return to_xlsx(rows, title)
    raise ValueError(fmt)


def _export_extra_fields(ids: list[int]) -> dict[int, dict]:
    """PMID, source, URL and keywords of the articles: the relevance query does not carry them."""
    if not ids:
        return {}
    out: dict[int, dict] = {}
    try:
        with engine.connect() as conn:
            for r in conn.execute(text(
                "SELECT id, pmid, source, url, keywords FROM literature_document WHERE id = ANY(:ids)"
            ), {"ids": [int(i) for i in ids]}).mappings():
                out[int(r["id"])] = dict(r)
    except Exception as _e:                                  # a database without `url`, say
        logger.warning(f"export extra fields: {_e}")
    return out


def _slug(s: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", (s or "").strip()).strip("-").lower()
    return s[:60] or "scenario"


def _normalize_format(fmt: str) -> str:
    """Le format demandé, ou 400. `bib` est un alias courant de `bibtex`."""
    fmt = (fmt or "csv").lower()
    if fmt == "bib":
        fmt = "bibtex"
    if fmt not in EXPORT_FORMATS:
        raise HTTPException(status_code=400,
                            detail=f"format must be one of {', '.join(EXPORT_FORMATS)}")
    return fmt


# Colonnes lues pour un export par identifiants. Mêmes noms que
# `_get_above_threshold_articles`, pour que `export_rows` ne voie aucune différence
# entre un sous-ensemble et le corpus pertinent.
_BY_IDS_SQL = """
    SELECT d.id, d.title, d.year, d.journal, d.authors, d.doi, d.study_design,
           d.citation_count, d.quality_score, d.abstract, d.pico_json,
           COALESCE(ars.screening_status, d.screening_status) AS screening_status,
           ars.similarity_score
    FROM literature_document d
    JOIN article_scenarios ars ON ars.document_id = d.id AND ars.scenario_id = :sid
    WHERE d.id = ANY(:ids) AND d.is_duplicate IS NOT TRUE
    ORDER BY (COALESCE(ars.screening_status, d.screening_status) = 'included') DESC,
             COALESCE(ars.rerank_score, ars.similarity_score, 0) DESC NULLS LAST,
             d.citation_count DESC NULLS LAST, d.id
"""


def articles_by_ids(scenario_id: str, ids: list[int]) -> list[dict]:
    """Les articles de CE scénario parmi `ids`, dans l'ordre de pertinence de l'app.

    Bornée au scénario par le JOIN : un identifiant qui ne lui appartient pas est
    simplement absent du résultat, jamais exporté. C'est ce qui permet d'ouvrir un
    export générique par identifiants sans en faire une fuite du corpus entier."""
    ids = [int(i) for i in ids if str(i).strip()]
    if not ids:
        return []
    with engine.connect() as conn:
        rows = conn.execute(text(_BY_IDS_SQL), {"sid": scenario_id, "ids": ids}).mappings().all()
    return [dict(r) for r in rows]


def articles_export_response(scenario_id: str, fmt: str, articles: list[dict],
                             include_abstract: bool, *, subset: str, subset_label: str,
                             coverage: str = "", extra_meta: dict | None = None) -> Response:
    """Rend N'IMPORTE QUEL sous-ensemble d'articles dans les six formats.

    Les formateurs étaient déjà purs ; seule la SÉLECTION des lignes était câblée sur le
    corpus pertinent. Ce noyau la prend en paramètre, de sorte qu'un cluster, un concept,
    une réponse du RAG ou une sélection à la main produisent exactement le même fichier,
    avec les mêmes colonnes.

    `coverage` : ce que le sous-ensemble couvre ET ce qu'il ne couvre pas (un cluster est
    tiré d'une projection plafonnée, pas du corpus entier). Il part dans les métadonnées
    du fichier, pas seulement dans l'interface : un export qui circule seul doit porter
    ses propres limites."""
    fmt = _normalize_format(fmt)
    row = _get_user_scenario_or_404(scenario_id)
    extra = _export_extra_fields([a["id"] for a in articles])
    for a in articles:
        a.update({k: v for k, v in (extra.get(int(a["id"])) or {}).items() if k != "id"})
    rows = export_rows(articles, include_abstract=include_abstract)
    title = str(row.get("name") or scenario_id)
    meta = {"scenario_id": scenario_id, "scenario": title, "query": row.get("query"),
            "n_articles": len(rows), "format": fmt,
            "subset": subset, "subset_label": subset_label}
    if coverage:
        meta["coverage"] = coverage
    meta.update(extra_meta or {})
    doc_title = title if subset == "relevant" else f"{title} - {subset_label}"
    body = render_export(fmt, rows, doc_title, meta)
    filename = f"{_slug(title)}_{_slug(subset_label)}_{len(rows)}.{_EXTENSIONS[fmt]}"
    return Response(content=body, media_type=_CONTENT_TYPES[fmt],
                    headers={"Content-Disposition": f'attachment; filename="{filename}"',
                             "X-Article-Count": str(len(rows)),
                             "X-Export-Subset": subset})


def relevant_articles_export(scenario_id: str, fmt: str, threshold: float | None,
                             include_abstract: bool) -> Response:
    articles = _get_above_threshold_articles(scenario_id, threshold=threshold)
    return articles_export_response(
        scenario_id, fmt, articles, include_abstract,
        subset="relevant", subset_label="relevant-articles",
        coverage=("Tous les articles pertinents du scénario : au-dessus du seuil de "
                  "similarité ou inclus par un relecteur, jamais les exclus."))


@app.get("/user-scenarios/{scenario_id}/relevant/export")
def export_user_scenario_relevant(
    scenario_id: str,
    format: str = Query("csv"),
    threshold: float | None = None,
    include_abstract: bool = True,
) -> Response:
    """The relevant articles of a scenario as a file: csv, xlsx, ris (Zotero, EndNote,
    Mendeley), bibtex, json or md. Same set and same order as the Corpus tab."""
    return relevant_articles_export(scenario_id, format, threshold, include_abstract)


@app.get("/gesica/scenarios/{scenario_id}/relevant/export")
def export_gesica_scenario_relevant(
    scenario_id: str,
    format: str = Query("csv"),
    threshold: float | None = None,
    include_abstract: bool = True,
) -> Response:
    """Same export for the built-in scenarios."""
    return relevant_articles_export(scenario_id, format, threshold, include_abstract)


# ─── Exports de SOUS-ENSEMBLES du corpus ─────────────────────────────────────
# Partout où l'application découpe le corpus (clusters, carte des concepts, réponse du
# RAG, sélection à la main), on doit pouvoir sortir la liste des articles correspondante,
# dans les mêmes six formats et avec les mêmes colonnes que l'export des pertinents. Les
# formateurs ne changent pas : seule la sélection des lignes diffère.

@app.get("/user-scenarios/{scenario_id}/clusters/{cluster_id}/export")
def export_scenario_cluster(
    scenario_id: str,
    cluster_id: int,
    format: str = Query("csv"),
    lang: str | None = Query(None),
    include_abstract: bool = True,
) -> Response:
    """Les articles d'UN cluster, dans le format demandé.

    Le cache de clustering conserve TOUS les points avec leur identifiant de document
    (seule la charge utile servie est sous-échantillonnée pour l'affichage), donc
    l'export d'un cluster est complet POUR CE CLUSTER.

    En revanche le clustering lui-même ne porte que sur les `CLUSTER_MAX_DOCS` articles
    les plus pertinents : c'est une projection bornée par la mémoire, pas une extraction.
    Le fichier le dit dans ses métadonnées, pour qu'un export qui circule seul ne se lise
    pas comme « tous les articles de ce thème »."""
    from .clustering import CLUSTER_MAX_DOCS, _load_viz_cache

    _get_user_scenario_or_404(scenario_id)
    cache = _load_viz_cache(scenario_id, "clustering_json")
    if not cache or not cache.get("clusters"):
        raise HTTPException(status_code=404,
                            detail="Aucun clustering en cache pour ce scénario : ouvrez "
                                   "l'onglet Clusters pour le calculer, puis réessayez.")
    match = next((c for c in cache["clusters"]
                  if isinstance(c, dict) and int(c.get("cluster_id", -99)) == int(cluster_id)), None)
    if match is None:
        _known = sorted(int(c.get("cluster_id")) for c in cache["clusters"]
                        if isinstance(c, dict) and c.get("cluster_id") is not None)
        raise HTTPException(status_code=404,
                            detail=f"Cluster {cluster_id} inconnu (disponibles : {_known}).")
    ids = [int(p["id"]) for p in (match.get("points") or [])
           if isinstance(p, dict) and p.get("id") is not None]
    articles = articles_by_ids(scenario_id, ids)
    _name = str(match.get("cluster_name") or f"cluster-{cluster_id}")
    _n_clustered = int(cache.get("n_docs") or 0)
    _n_eligible = int(cache.get("n_docs_total") or _n_clustered)
    coverage = (
        f"Articles du cluster « {_name} » ({len(articles)}). Le clustering porte sur "
        f"{_n_clustered} articles sur {_n_eligible} éligibles : c'est une projection "
        f"plafonnée à CLUSTER_MAX_DOCS={CLUSTER_MAX_DOCS} (les plus pertinents), pas le "
        f"corpus entier. Pour la liste complète des articles pertinents, utilisez "
        f"l'export du corpus."
    )
    return articles_export_response(
        scenario_id, format, articles, include_abstract,
        subset="cluster", subset_label=_slug(_name) or f"cluster-{cluster_id}",
        coverage=coverage,
        extra_meta={"cluster_id": int(cluster_id), "cluster_name": _name,
                    "clustering_lang": cache.get("lang") or lang,
                    "n_docs_clustered": _n_clustered, "n_docs_eligible": _n_eligible,
                    "top_words": match.get("top_words") or [],
                    "cluster_summary": match.get("summary") or ""})


@app.get("/user-scenarios/{scenario_id}/concepts/export")
def export_scenario_concept_subset(
    scenario_id: str,
    concepts: str = Query(..., description="type:label, séparés par « | ». Ex: pathogen:dengue virus|vector:Aedes albopictus"),
    mode: str = Query("any", pattern="^(any|all)$"),
    format: str = Query("csv"),
    include_abstract: bool = True,
) -> Response:
    """Les articles derrière un concept, ou derrière une SÉLECTION de concepts.

    `mode=any` : les articles citant AU MOINS UN des concepts (l'union, ce que montre une
    carte filtrée). `mode=all` : ceux qui les citent TOUS (l'intersection, ce qu'on veut
    pour « les articles où ce pathogène ET ce vecteur apparaissent ensemble »).

    La carte est RECALCULÉE sans le plafond d'affichage : les noeuds servis à l'interface
    ne portent que leurs 40 premiers articles, et hériter de ce plafond ici produirait un
    fichier tronqué sans le dire. Les concepts sont désignés par `type:label` (le libellé
    canonique anglais) et non par leur identifiant de noeud, qui n'est qu'une position
    dans un calcul donné et change d'un recalcul à l'autre."""
    from .knowledge_graph import _build_concept_graph, _concept_rows

    _get_user_scenario_or_404(scenario_id)
    wanted: list[tuple[str, str]] = []
    for part in (concepts or "").split("|"):
        part = part.strip()
        if not part:
            continue
        if ":" not in part:
            raise HTTPException(status_code=400,
                                detail=f"« {part} » doit s'écrire type:label (ex. pathogen:dengue virus).")
        _t, _lab = part.split(":", 1)
        wanted.append((_t.strip().lower(), _lab.strip()))
    if not wanted:
        raise HTTPException(status_code=400, detail="Aucun concept demandé.")

    rows, n_total = _concept_rows(scenario_id)
    graph = _build_concept_graph(rows, n_total=n_total, full_articles=True)
    by_key = {(str(n["type"]).lower(), str((n.get("label") or {}).get("en") or "").strip().lower()): n
              for n in graph.get("nodes") or []}
    sets: list[set[int]] = []
    labels: list[str] = []
    missing: list[str] = []
    for t, lab in wanted:
        node = by_key.get((t, lab.lower()))
        if node is None:
            missing.append(f"{t}:{lab}")
            continue
        sets.append({int(i) for i in (node.get("articles") or [])})
        labels.append(str((node.get("label") or {}).get("en") or lab))
    if missing:
        raise HTTPException(
            status_code=404,
            detail=f"Concept(s) absent(s) de la carte : {', '.join(missing)}. "
                   f"Les libellés sont ceux de la carte (anglais canonique).")
    ids = set.union(*sets) if mode == "any" else set.intersection(*sets)
    articles = articles_by_ids(scenario_id, sorted(ids))
    _joiner = " OU " if mode == "any" else " ET "
    _human = _joiner.join(labels)
    coverage = (
        f"Articles citant {_human} ({len(articles)}). Calculé sur la TOTALITÉ des "
        f"{graph.get('n_with_concepts', 0)} articles pertinents porteurs de concepts "
        f"(sur {graph.get('n_total', n_total)}), sans le plafond d'affichage de la carte."
    )
    return articles_export_response(
        scenario_id, format, articles, include_abstract,
        subset="concepts", subset_label=_slug("-".join(labels)) or "concepts",
        coverage=coverage,
        extra_meta={"concepts": labels, "mode": mode,
                    "n_with_concepts": graph.get("n_with_concepts"),
                    "n_total": graph.get("n_total", n_total)})


@app.get("/user-scenarios/{scenario_id}/articles/export")
def export_scenario_article_ids(
    scenario_id: str,
    ids: str = Query(..., description="Identifiants d'articles séparés par des virgules."),
    format: str = Query("csv"),
    label: str = Query("selection", description="Nom du sous-ensemble, pour le fichier."),
    include_abstract: bool = True,
) -> Response:
    """Export d'une LISTE EXPLICITE d'articles de ce scénario.

    C'est l'export générique : il sert les sources d'une réponse du RAG (qui portent leur
    `document_id`), une carte des concepts filtrée à l'écran, une sélection faite à la
    main, et tout découpage à venir, sans qu'il faille un endpoint de plus à chaque fois.

    Borné au scénario : un identifiant qui ne lui appartient pas n'est pas exporté, il est
    signalé dans `X-Missing-Ids`. La réponse ne peut donc pas servir à extraire le corpus
    d'un autre scénario en devinant des identifiants."""
    try:
        wanted = [int(x) for x in (ids or "").replace(" ", "").split(",") if x]
    except ValueError:
        raise HTTPException(status_code=400, detail="`ids` doit être une liste d'entiers séparés par des virgules.")
    if not wanted:
        raise HTTPException(status_code=400, detail="Aucun identifiant fourni.")
    if len(wanted) > 20000:
        raise HTTPException(status_code=400, detail="Trop d'identifiants (maximum 20000).")
    _get_user_scenario_or_404(scenario_id)
    articles = articles_by_ids(scenario_id, wanted)
    _found = {int(a["id"]) for a in articles}
    _missing = [i for i in wanted if i not in _found]
    coverage = (
        f"Sélection explicite de {len(articles)} article(s) sur {len(wanted)} demandé(s)."
        + (f" {len(_missing)} identifiant(s) n'appartiennent pas à ce scénario et ont été "
           f"écartés." if _missing else "")
    )
    resp = articles_export_response(
        scenario_id, format, articles, include_abstract,
        subset="selection", subset_label=_slug(label) or "selection",
        coverage=coverage,
        extra_meta={"requested": len(wanted), "returned": len(articles),
                    "missing_ids": _missing[:100]})
    if _missing:
        resp.headers["X-Missing-Ids"] = str(len(_missing))
    return resp
