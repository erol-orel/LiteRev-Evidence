"""Structured extraction in the shape of the T2.4 / T4.5 data extraction template.

One row per observation (a covariate of a study population), the way the reviewers' Excel
template wants it: a REF row per paper, then one long-format sheet per covariate group
(human susceptibility, human exposure, environment, animal or reservoir, vector). Every
observation carries the exact quote it comes from, the page or section, whether it came
from a table, a figure or the text, and whether the quote was found in the text that the
model was given. A paper that does not report something says so ("coverage" flags) instead
of staying silent.

House rule (CLAUDE.md): every extraction draws on ALL the relevant articles. This is the
MAP half: per-article facts are extracted once and cached on the article row
(`extraction_json`), so the cost is one-time and incremental; a paper is re-read only when
the full text arrives after an abstract-only pass, or when the extraction version moves.
The REDUCE half (counts and coverage over the whole relevant subset, in SQL) is the next
step. `EXTRACTION_MAX_ARTICLES` defaults to zero, meaning no limit; a positive value is an
operational fallback for a day when the LLM budget must be held.

The per-article TEXT is capped (`EXTRACTION_MAX_CHARS`): that is the size of one prompt,
not a sample of the corpus, and the extraction records when it had to cut.
"""
from __future__ import annotations

import csv
import functools
import hashlib
import io
import json
import os
import re
import threading
import time
import unicodedata
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from typing import Any

from fastapi import Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import text

from .codebook import annotate, get_codebook, get_index, normalise, sheet_key, vocabulary_prompt
from .core import (_env_int, _is_openai_quota_error, _openai_in_cooldown, _trip_openai_cooldown, app,
                   engine, logger, require_api_key)
from .extraction_review import load_reviews, overlay, overlay_annotated
from .geography import geography, get_nuts_index, resolve_location
from .digest import _rows
from .scenario_store import (_get_scenario_threshold, _get_user_scenario_or_404,
                             relevant_gate_sql, screening_status_sql)
from .schema_boot import _exec_ddl_isolated
from llm_usage import LLMTruncated, ProviderWatch, is_transient_llm_error
from llm_usage import json_content as _json_content
from llm_usage import model_for as _model

#: Bump to make every cached extraction eligible for a new pass.
EXTRACTION_VERSION = 1
EXTRACTION_MAX_ARTICLES = _env_int("EXTRACTION_MAX_ARTICLES", 0, 0)
#: Characters of a paper's text read per article (the title, then the text). 150,000 is about
#: 37,000 tokens, under a cent per paper at the bulk price, and holds a whole paper in nearly
#: every case. It was 60,000, minus the chunk overlaps read twice: 149 papers of HPAI_last were
#: cut, at the end of the paper, where PDF tables and supplementary tables sit.
EXTRACTION_MAX_CHARS = _env_int("EXTRACTION_MAX_CHARS", 150000, 4000)
#: The limit an extraction was cut at when it does not say (it was made before `max_chars` was
#: recorded): such a cut paper is read again once the limit is above it.
_LEGACY_MAX_CHARS = 60000
_EXTRACTION_WORKERS = _env_int("EXTRACTION_WORKERS", 4, 1)
#: Failed attempts after which an article is left alone (a deterministic failure would
#: otherwise be paid for again at every run). A success resets the counter. A transient
#: failure (an outage, a refused call, a quota) is never charged.
_MAX_ATTEMPTS = 3
_MAX_OBSERVATIONS = 120
_MIN_ABSTRACT_CHARS = 80
#: The answer's ceiling, in tokens. The prompt allows 120 observations with quotes of up to
#: 300 characters, which needs about 18,000 to 21,000 tokens. At 8,000 a rich paper (a
#: meta-analysis, a serology table, a KAP survey) had its answer cut mid-JSON, failed
#: identically at every attempt (temperature 0, fixed seed) and was given up: on HPAI_last no
#: extracted article held more than 72 rows. Output is billed only as written.
_EXTRACTION_MAX_TOKENS = _env_int("EXTRACTION_MAX_TOKENS", 24000, 4000)
#: What is asked instead when the API refuses the ceiling as too large for the model.
_FALLBACK_MAX_TOKENS = 8000
#: A 24,000-token answer takes minutes to write: the client waits for it.
_LLM_TIMEOUT_S = 600.0
#: When the answer is cut, ONE more call asks for a condensed answer that fits.
_CONDENSED_MAX_OBSERVATIONS = 50
_CONDENSED_RULE = (
    "\n\nTHIS PAPER HOLDS MORE DATA THAN ONE ANSWER CAN CARRY: a first answer was cut off. "
    "Return at most " + str(_CONDENSED_MAX_OBSERVATIONS) + " observations, each quote at most "
    "150 characters. Keep first every row that gives counts (a number of cases or positives "
    "out of a number of people or animals), on every sheet: human, environment, animal and "
    "vector alike. Then the sex, gender, age, occupation, knowledge-attitude-practice and "
    "protective-equipment results. Then the rest."
)

SHEETS = ("human_susc", "human_exp", "env", "animal", "vector")
SOURCE_KINDS = ("table", "figure", "text")
#: What a paper can report on, true only when it gives data on it. A False is information:
#: "this paper has no sex-disaggregated data" is a result for the gap map, not a blank.
COVERAGE_KEYS = ("sex_gender", "age", "occupation", "kap_risk_perception", "ppe",
                 "vaccination", "human_testing", "animal_host", "environment", "vector")

# ── The template, verbatim ───────────────────────────────────────────────────
# Column titles exactly as the reviewers' workbook spells them (trailing and doubled
# spaces included), so that what we export pastes into their file without remapping.
REF_COLUMNS = (
    ("ID", "id"), ("first author name", "first_author"), ("reference", "reference"),
    ("year of publication", "year"), ("brief study description", "description"),
    ("Article type", "article_type"),
    ("Starting date of the study (DD/MM/YYYY)", "study_start"),
    ("Ending date of the study (DD/MM/YYYY)", "study_end"),
    ("geographic location and area", "location"),
    ("study area NUTS level 1 ", "nuts1"), ("study area NUTS level 2 ", "nuts2"),
    ("study area NUTS level 3", "nuts3"),
    ("Notes (space and time additional explanation)", "notes_geo"),
    ("Number of exposed (population at risk)", "risk_pop"),
    ("Number  of positive (cases)", "positive"),
    ("Proportion % of positive (cases)", "percent_positive"),
    ("Mathematical model (Y/N)", "math_model"), ("Model type", "model_type"),
    ("Exclusion", "exclusion"),
)
_HUM = (("ID", "id"), ("TRANSMISSION MODE", "transmission_mode"), ("DISEASE", "disease"),
        ("COV_HUM_GROUP", "group"), ("COVARIATE_hum", "covariate"), ("Value_hum", "value"),
        ("descr_hum", "descr"), ("notes_hum", "notes"), ("N_cases_hum", "n_cases"),
        ("Pop_risk_Hum", "pop_risk"), ("Variable original name", "original_name"),
        ("Page/section", "page_section"), ("Extraction from Table/Figure/text", "source_kind"))
SHEET_COLUMNS = {
    "human_susc": ("HUMAN_COV_SUSC", _HUM),
    "human_exp": ("HUMAN_COV_EXP", _HUM),
    "env": ("ENV_COV", (
        ("ID", "id"), ("TRANSMISSION MODE", "transmission_mode"), ("DISEASE", "disease"),
        ("COVARIATE_env", "covariate"), ("Value_env", "value"), ("descr_env", "descr"),
        ("notes_env", "notes"), ("N_cases_env", "n_cases"), ("Pop_risk_env", "pop_risk"),
        ("Extraction from Table/Figure/text", "source_kind"))),
    "animal": ("ANIMALorRESERVOIR_COV", (
        ("ID", "id"), ("TRANSMISSION MODE", "transmission_mode"), ("DISEASE", "disease"),
        ("COV_ANIM_GROUP", "group"), ("COVARIATE_res", "covariate"), ("Value_res", "value"),
        ("descr_res", "descr"), ("notes_res", "notes"), ("N_cases_res", "n_cases"),
        ("Pop_risk_res", "pop_risk"), ("Extraction from Table/Figure/text", "source_kind"))),
    "vector": ("VECTOR_COV", (
        ("ID", "id"), ("TRANSMISSION MODE", "transmission_mode"), ("DISEASE", "disease"),
        ("COV_VEC_GROUP", "group"), ("COVARIATE_vec", "covariate"), ("Value_vec", "value"),
        ("descr_vec", "descr"), ("notes_vec", "notes"), ("N_cases_vec", "n_cases"),
        ("Pop_risk_vec", "pop_risk"), ("Extraction from Table/Figure/text", "source_kind"))),
}
#: Added to the right of the template columns, so the template stays intact and a reviewer
#: can still check each value against the paper.
REVIEW_COLUMNS = (("Page/section", "page_section"), ("Codebook label", "label_path"),
                  ("Review status", "review_status"), ("Reviewed by", "reviewed_by"), ("Quote", "quote"),
                  ("Quote found in text", "quote_verified"), ("Extracted from", "source"))
#: On the right of REF: what was read of each paper, and whether the answer was condensed.
REF_READ_COLUMNS = (("Extracted from", "source"), ("Characters read", "chars_read"),
                    ("Text cut at the size limit (Y/N)", "text_cut"), ("Condensed answer (Y/N)", "condensed"))


def _ensure_extraction_columns() -> None:
    stmts = [
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS extraction_json JSONB",
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS extraction_at TIMESTAMP",
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS extraction_attempts INTEGER DEFAULT 0",
        # Paramètres épidémiologiques extraits PAR ARTICLE. La règle de la maison : le
        # travail par article est mis en cache sur sa ligne (comme pico_json et
        # concepts_json), et seule la réduction est refaite. Cette extraction ne cachait
        # rien et renvoyait au modèle les mêmes résumés à chaque régénération.
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS epi_params_json JSONB",
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS epi_params_at TIMESTAMP",
        # Le lien d'accès ouvert rendu par Unpaywall. Il était écrit dans `url` en
        # posant `has_fulltext = true`, sans stocker une ligne de texte : le panneau
        # annonçait « texte intégral récupéré » pour une adresse.
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS oa_url TEXT",
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS oa_url_found_at TIMESTAMP",
        # Tentatives d'annotation des concepts. Sans ce compteur, un article qu'aucun
        # passage ne sait annoter restait « manquant » et relançait un passage payant à
        # chaque affichage de la carte des concepts, jusqu'à quarante fois par page.
        "ALTER TABLE literature_document ADD COLUMN IF NOT EXISTS concepts_attempts INTEGER DEFAULT 0",
    ]
    _exec_ddl_isolated(stmts, "_ensure_extraction_columns")


try:
    _ensure_extraction_columns()
except Exception as _e:                                       # noqa: BLE001 - never blocks startup
    logger.warning(f"_ensure_extraction_columns: {_e}")


# ─────────────────────────────────────────────────────────────────────────────
# The prompt
# ─────────────────────────────────────────────────────────────────────────────
_EXTRACTION_SYSTEM = (
    "You extract structured data from ONE scientific paper for a systematic review on "
    "emerging infectious diseases (One Health: humans, animals, environment, vectors). "
    "Report ONLY what the text states: never infer, never estimate, never carry a value over "
    "from another paper or another disease. If the paper does not report something, leave it "
    "out. Return ONLY JSON of this shape:\n"
    '{"ref": {"description": <one or two sentences: study type, population, setting>, '
    '"article_type": <research | short communication | outbreak report | review | other>, '
    '"study_start": <string as written or null>, "study_end": <string as written or null>, '
    '"location": <place name as written, or null>, "notes_geo": <string or null>, '
    '"risk_pop": <number or null>, "positive": <number or null>, '
    '"percent_positive": <number or null>, "math_model": <true|false>, "model_type": '
    '<e.g. SEIR, or null>},\n'
    ' "coverage": {"sex_gender": <bool>, "age": <bool>, "occupation": <bool>, '
    '"kap_risk_perception": <bool>, "ppe": <bool>, "vaccination": <bool>, "human_testing": '
    '<bool>, "animal_host": <bool>, "environment": <bool>, "vector": <bool>},\n'
    ' "observations": [{"sheet": <one of the five below>, "transmission_mode": <string or '
    'null>, "disease": <string>, "group": <string>, "covariate": <string>, "value": '
    '<number or null>, "descr": <string>, "notes": <string or null>, "n_cases": <number or '
    'null>, "pop_risk": <number or null>, "original_name": <the label used in the paper>, '
    '"page_section": <section name, table or figure number, or page>, "source_kind": '
    '<table | figure | text>, "quote": <verbatim excerpt of the text>}]}\n\n'
    "SHEETS. human_susc: host factors that change susceptibility or outcome (sex, gender, "
    "age, comorbidity, immune or vaccination status as a host factor). human_exp: factors "
    "that change exposure (occupation, type of contact, behaviour, knowledge, attitudes and "
    "practices, risk perception, protective equipment use). env: environmental covariates "
    "(temperature, humidity, biosecurity level, setting, wild-bird contact, persistence of "
    "the pathogen). animal: animal hosts and reservoirs (species, number tested, positive, "
    "serology, clinical status). vector: vectors (species, density, infection rate).\n\n"
    "RULES. One row per covariate per group: a sex-disaggregated result gives one row for "
    "males and one for females, with the group 'sex' and the covariate 'male' or 'female'. "
    "'value' is a single number as reported (proportion, odds ratio, mean, count); put the "
    "unit, the confidence interval and the statistic name in 'descr'. 'n_cases' is the number "
    "of cases or positives in that group and 'pop_risk' the number of people or animals in it; "
    "give them only when the paper does. 'quote' must be copied exactly from the text, at most "
    "300 characters, and must contain the figures you report. Scientific names in Latin for "
    "animals and vectors. At most " + str(_MAX_OBSERVATIONS) + " observations: keep the most "
    "informative, and always keep every sex, gender, age, occupation, knowledge-attitude-"
    "practice and protective-equipment result. 'coverage' flags are true only when the paper "
    "gives data on that item, false otherwise. Never write the em dash character; use a comma "
    "or a hyphen. Return ONLY the JSON."
)


#: Which prompt produced an extraction, for the reproducibility record: any edit to the prompt
#: changes it, so two extractions of the same corpus can be told apart.
PROMPT_SHA = hashlib.sha256(_EXTRACTION_SYSTEM.encode("utf-8")).hexdigest()[:10]


# ─────────────────────────────────────────────────────────────────────────────
# Pure parsing and checking
# ─────────────────────────────────────────────────────────────────────────────
_NUM_RE = re.compile(r"^-?\d+(?:[.,]\d+)?$")


def _to_number(v: Any) -> int | float | None:
    """A number from what the model returned, or None. Integral values stay integers."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        f = float(v)
    elif isinstance(v, str) and _NUM_RE.match(v.strip().replace(" ", "")):
        f = float(v.strip().replace(" ", "").replace(",", "."))
    else:
        return None
    if f != f or f in (float("inf"), float("-inf")):
        return None
    return int(f) if f.is_integer() and abs(f) < 1e15 else f


def _clean_str(v: Any, limit: int) -> str | None:
    if v is None or isinstance(v, (dict, list)):
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s[:limit] if s else None


#: Curly quotes, primes and every dash or minus, as their plain ASCII forms.
_PLAIN_PUNCT = str.maketrans({
    **{chr(c): "'" for c in (0x2018, 0x2019, 0x201A, 0x201B, 0x2032, 0x00B4)},
    **{chr(c): '"' for c in (0x201C, 0x201D, 0x201E, 0x201F, 0x2033, 0x00AB, 0x00BB)},
    **{chr(c): "-" for c in (0x2010, 0x2011, 0x2012, 0x2013, 0x2014, 0x2015, 0x2212, 0xFE58, 0xFE63, 0xFF0D)},
})
_SOFT_HYPHEN = re.compile(chr(0x00AD) + r"\s*")
#: A word cut at a line break: "trans-" then "mission". The full-text chunker collapses every
#: line break to a space, so a hyphen followed by any whitespace counts.
_HYPHEN_BREAK = re.compile(r"(?<=\w)-\s+(?=\w)")
#: Combining marks, i.e. the accents of a decomposed letter: one regex pass, not a loop per
#: character (a 150,000-character paper is folded in milliseconds).
_COMBINING = re.compile("[" + "".join(f"{chr(a)}-{chr(b)}" for a, b in (
    (0x0300, 0x036F), (0x1AB0, 0x1AFF), (0x1DC0, 0x1DFF), (0x20D0, 0x20FF), (0xFE20, 0xFE2F))) + "]")


_NOT_A_WORD = re.compile(r"[^\w\s%.]|_")


def _plain(s: str) -> str:
    """NFKC (the ligature "fi" of a PDF is two letters again), curly quotes and dashes plain."""
    return unicodedata.normalize("NFKC", s or "").translate(_PLAIN_PUNCT)


def _fold(t: str) -> str:
    """A `_plain` text as words and figures only: accents off, soft hyphens out, case folded,
    and anything that is neither a letter, a digit, "%" nor "." as one space."""
    t = _COMBINING.sub("", unicodedata.normalize("NFD", t))
    t = _SOFT_HYPHEN.sub("", t).casefold()
    return " ".join(_NOT_A_WORD.sub(" ", t).split())


def _norm_text(s: str) -> str:
    return _fold(_plain(s))


@functools.lru_cache(maxsize=8)
def _haystack(source_text: str, joined: bool) -> str:
    """The text as the quote check reads it: as written, or with the words hyphenated across a
    line break joined (a real compound, "COVID-19", may sit at a line break too, so both are
    tried). One text is read for every row of an article: it is folded once, not per row."""
    plain = _plain(source_text)
    return _fold(_HYPHEN_BREAK.sub("", plain) if joined else plain)


def quote_in_text(quote: str | None, source_text: str) -> bool:
    """Whether the model's quote is really in the text it was given.

    Whitespace, case, accents and punctuation are ignored, and what a PDF does to text is
    undone on both sides (ligatures, soft hyphens, words split at a line break, curly quotes),
    but the words and figures must be there in the same order. This is the cheap check that
    catches an invented excerpt; it does not prove the value is the right reading of it."""
    q = _norm_text(quote or "")
    if not q:
        return False
    return q in _haystack(source_text or "", False) or q in _haystack(source_text or "", True)


def parse_extraction(raw: Any, source_text: str) -> dict[str, Any]:
    """The model's JSON as a clean, bounded record. Pure.

    A sheet written another way ("Human_Susc", "environment", "ENV_COV") is read as the sheet
    it names. Rows with no recognisable sheet or no covariate are dropped, and COUNTED, with
    the reason (`n_rows_dropped`, `rows_dropped`): they used to vanish without a trace. A value
    that is not a number is kept as text in `descr` rather than lost. `quote_verified` says
    whether the quote was found in `source_text`; an unverified row is KEPT and flagged, so the
    reviewer sees it instead of the model's hallucination being hidden by a silent filter."""
    if not isinstance(raw, dict):
        raise ValueError("extraction is not a JSON object")
    ref_in = raw.get("ref") if isinstance(raw.get("ref"), dict) else {}
    ref = {
        "description": _clean_str(ref_in.get("description"), 500),
        "article_type": _clean_str(ref_in.get("article_type"), 60),
        "study_start": _clean_str(ref_in.get("study_start"), 40),
        "study_end": _clean_str(ref_in.get("study_end"), 40),
        "location": _clean_str(ref_in.get("location"), 200),
        "notes_geo": _clean_str(ref_in.get("notes_geo"), 400),
        "risk_pop": _to_number(ref_in.get("risk_pop")),
        "positive": _to_number(ref_in.get("positive")),
        "percent_positive": _to_number(ref_in.get("percent_positive")),
        "math_model": bool(ref_in.get("math_model")) if isinstance(ref_in.get("math_model"), bool) else False,
        "model_type": _clean_str(ref_in.get("model_type"), 80),
    }
    cov_in = raw.get("coverage") if isinstance(raw.get("coverage"), dict) else {}
    coverage = {k: cov_in.get(k) is True for k in COVERAGE_KEYS}

    observations: list[dict[str, Any]] = []
    dropped: Counter = Counter()
    for o in (raw.get("observations") or []) if isinstance(raw.get("observations"), list) else []:
        if not isinstance(o, dict):
            dropped["not_an_object"] += 1
            continue
        sheet = sheet_key(_clean_str(o.get("sheet"), 60))
        covariate = _clean_str(o.get("covariate"), 200)
        if sheet not in SHEETS:
            dropped["unknown_sheet"] += 1
            continue
        if not covariate:
            dropped["no_covariate"] += 1
            continue
        if len(observations) >= _MAX_OBSERVATIONS:
            dropped["over_limit"] += 1
            continue
        descr = _clean_str(o.get("descr"), 400)
        value = _to_number(o.get("value"))
        if value is None and o.get("value") not in (None, ""):
            as_text = _clean_str(o.get("value"), 120)
            if as_text:
                descr = f"{as_text}. {descr}" if descr else as_text
        kind = (_clean_str(o.get("source_kind"), 12) or "").lower()
        quote = _clean_str(o.get("quote"), 300)
        observations.append({
            "sheet": sheet,
            "transmission_mode": _clean_str(o.get("transmission_mode"), 80),
            "disease": _clean_str(o.get("disease"), 120),
            "group": _clean_str(o.get("group"), 120),
            "covariate": covariate,
            "value": value,
            "descr": descr,
            "notes": _clean_str(o.get("notes"), 400),
            "n_cases": _to_number(o.get("n_cases")),
            "pop_risk": _to_number(o.get("pop_risk")),
            "original_name": _clean_str(o.get("original_name"), 200),
            "page_section": _clean_str(o.get("page_section"), 120),
            "source_kind": kind if kind in SOURCE_KINDS else None,
            "quote": quote,
            "quote_verified": quote_in_text(quote, source_text),
        })
    return {"ref": ref, "coverage": coverage, "observations": observations,
            "n_rows_dropped": sum(dropped.values()), "rows_dropped": dict(dropped)}


# ─────────────────────────────────────────────────────────────────────────────
# Reading one article
# ─────────────────────────────────────────────────────────────────────────────
#: The chunker repeats up to 400 characters of whole sentences at the head of the next chunk
#: (pipeline `_chunk_text_ft`); older imports repeated 300. A repeat is looked for up to this
#: length, and only from this one (a shorter match could be chance).
_MAX_CHUNK_OVERLAP = 600
_MIN_CHUNK_OVERLAP = 20


def _chunk_overlap(prev: str, nxt: str) -> int:
    """Length of the longest end of `prev` that `nxt` starts with (0 when none is long enough)."""
    n = min(len(prev), len(nxt), _MAX_CHUNK_OVERLAP)
    if n < _MIN_CHUNK_OVERLAP:
        return 0
    tail, probe = prev[-n:], nxt[:_MIN_CHUNK_OVERLAP]
    i = tail.find(probe)
    while i != -1:                                    # the first hit is the longest overlap
        if nxt.startswith(tail[i:]):
            return n - i
        i = tail.find(probe, i + 1)
    return 0


def _join_fulltext_chunks(chunks: list[str]) -> str:
    """A full text from its chunks, in order, each overlap read once. The plain join repeated
    up to 400 characters at every boundary: about 6,000 of the 60,000 characters read were
    read twice, and the end of the paper was cut for them."""
    out: list[str] = []
    prev = ""
    for c in chunks:
        c = (c or "").strip()
        if not c:
            continue
        rest = c[_chunk_overlap(prev, c):].strip()
        if rest:
            out.append(rest)
        prev = c
    return "\n\n".join(out)


#: A reference list heading. Inline in the collapsed text of a chunk ("... interest. References
#: 1. Smith J ..."), so it must follow the end of a sentence, a line break or the start, unless
#: written in capitals; then a list must follow (see `_drop_references`).
_REF_HEADING = re.compile(
    r"(?:(?:^|(?<=[\n.!?:;)\]]))\s*(?:\d{1,2}\.?\s+)?"
    r"(?:References(?: and [Nn]otes| [Cc]ited)?|Bibliography|Literature [Cc]ited|Works [Cc]ited"
    r"|Reference [Ll]ist|R[ée]f[ée]rences(?: bibliographiques)?)"
    r"|\b(?:REFERENCES|BIBLIOGRAPHY|LITERATURE CITED|R[ÉE]F[ÉE]RENCES))\s*:?\s+(?=\S)")
_YEAR = re.compile(r"\b(?:19|20)\d\d\b")
#: What may follow a reference list and must still be read: tables and figure legends set
#: after the references (author manuscripts), supplementary material, appendices.
_AFTER_REFERENCES = re.compile(
    r"(?<![\w-])(?:Tables?\s+S?\d+|TABLES?\s+S?\d+|Supplementary\b|SUPPLEMENTARY\b|Supporting [Ii]nformation"
    r"|Appendix\b|APPENDIX\b|Annex\b|ANNEX\b|Figure [Ll]egends?|Fig(?:ure)?s?\.?\s+S?\d+[.:])")
_REFS_FROM = 0.6                                       # the heading is in the last 40% of the text
_MIN_REFS_CHARS = 300


def _drop_references(body: str) -> tuple[str, int]:
    """(text, characters dropped): the paper without its reference list, which carries no data
    for the extraction and was what the size limit left in when it cut tables at the end.

    Conservative: the heading must sit in the last 40% of the text and be followed by a list
    (three years within 800 characters); what comes after the list (tables, figure legends,
    supplementary material, appendices) is kept. Nothing found, nothing dropped."""
    for m in _REF_HEADING.finditer(body, int(len(body) * _REFS_FROM)):
        if len(_YEAR.findall(body[m.end():m.end() + 800])) < 3:
            continue
        after = _AFTER_REFERENCES.search(body, m.end())
        end = after.start() if after else len(body)
        if end - m.start() < _MIN_REFS_CHARS:
            continue
        return (body[:m.start()].rstrip() + ("\n\n" + body[end:] if after else "")), end - m.start()
    return body, 0


def _article_text_info(conn, row: dict) -> dict[str, Any]:
    """What the model reads for one article, and how much of the paper that is.

    `text` is the full text when the article has one (its chunks joined without their
    repeats, the reference list dropped, then cut at EXTRACTION_MAX_CHARS), else the abstract;
    the title always leads. `source` is what the reviewer must know: values extracted from an
    abstract are a fraction of what the paper reports (tables are not in abstracts).
    `chars_read` of `chars_total` characters were read under the limit `max_chars`;
    `truncated` says the text was cut."""
    title = (row.get("title") or "").strip()
    chunks = [r[0] for r in conn.execute(text("""
        SELECT content FROM document_chunk
        WHERE document_id = :id AND chunk_type IN ('fulltext_section', 'full_text')
        ORDER BY chunk_index, id
    """), {"id": row["id"]}).all()]
    full = _join_fulltext_chunks(chunks)
    limit = EXTRACTION_MAX_CHARS
    if full and len(full) > 500:
        body, refs = _drop_references(full)
        read = body[:limit]
        return {"text": f"{title}\n\n{read}", "source": "fulltext", "truncated": len(body) > limit,
                "chars_read": len(read), "chars_total": len(body), "max_chars": limit,
                "references_dropped": refs}
    abstract = (row.get("abstract") or "").strip()
    return {"text": f"{title}\n\n{abstract}", "source": "abstract", "truncated": False,
            "chars_read": len(abstract), "chars_total": len(abstract), "max_chars": limit,
            "references_dropped": 0}


def _article_text(conn, row: dict) -> tuple[str, str, bool]:
    """(text, source, truncated): `_article_text_info` without the counts."""
    info = _article_text_info(conn, row)
    return info["text"], info["source"], info["truncated"]


def _article_whole_text(conn, row: dict) -> str:
    """Everything the database holds of a paper: title, abstract and the whole full text,
    never cut. What a quote is checked against again (`reverify_quotes`): a quote found
    anywhere in the paper is not an invented one."""
    chunks = [r[0] for r in conn.execute(text("""
        SELECT content FROM document_chunk
        WHERE document_id = :id AND chunk_type IN ('fulltext_section', 'full_text')
        ORDER BY chunk_index, id
    """), {"id": row["id"]}).all()]
    return "\n\n".join(p for p in ((row.get("title") or "").strip(), (row.get("abstract") or "").strip(),
                                   _join_fulltext_chunks(chunks)) if p)


#: Models whose API refused the ceiling as too large, in this process: they are asked
#: `_FALLBACK_MAX_TOKENS` from then on, so the refusal is paid once, not once per article.
_CEILING_REFUSED: dict[str, int] = {}
_RE_CEILING = re.compile(r"max_completion_tokens|max_tokens|max(?:imum)? (?:number of )?(?:completion |output )?tokens",
                         re.I)


def _refuses_ceiling(exc: BaseException) -> bool:
    """A 400 invalid request about the token ceiling itself (too large for this model), and
    not the rename of `max_tokens` that llm_usage already repairs."""
    msg = str(exc)
    status = getattr(exc, "status_code", None)
    bad_request = status == 400 or type(exc).__name__ == "BadRequestError" or "error code: 400" in msg.lower()
    return bad_request and bool(_RE_CEILING.search(msg)) and "unsupported parameter" not in msg.lower()


def extract_article(client, row: dict, source_text: str, source: str, truncated: bool,
                    disease_hint: str | None = None, system: str | None = None,
                    text_info: dict | None = None) -> dict[str, Any]:
    """One LLM call for one article, parsed and checked. Raises on any failure, so the
    caller counts an attempt instead of caching an empty result as if it were an answer.
    `text_info` (from `_article_text_info`) says how much of the paper `source_text` is."""
    system = system or _EXTRACTION_SYSTEM
    payload = {"disease_of_interest": disease_hint or None, "text_is": source, "text": source_text}
    model = _model("bulk")
    ceiling = [min(_FALLBACK_MAX_TOKENS, _EXTRACTION_MAX_TOKENS) if model in _CEILING_REFUSED
               else _EXTRACTION_MAX_TOKENS]

    def _ask(system_prompt: str):
        def call(max_tokens: int):
            return client.chat.completions.create(
                model=model,
                messages=[{"role": "system", "content": system_prompt},
                          {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}],
                temperature=0, seed=42, max_tokens=max_tokens,
                response_format={"type": "json_object"},
            )
        try:
            return call(ceiling[0])
        except Exception as e:
            # A model whose answer cannot be that long: asked again at the old ceiling, in the
            # same attempt (nothing was read, nothing is charged), and remembered.
            if ceiling[0] <= _FALLBACK_MAX_TOKENS or not _refuses_ceiling(e):
                raise
            logger.warning(f"extraction {row.get('id')}: {model} refused max_tokens={ceiling[0]} ({e}); "
                           f"asking {_FALLBACK_MAX_TOKENS} from now on. Lower EXTRACTION_MAX_TOKENS.")
            _CEILING_REFUSED[model] = ceiling[0]
            ceiling[0] = _FALLBACK_MAX_TOKENS
            return call(ceiling[0])

    # An answer cut at its ceiling is no answer (json_content raises), and asking the same
    # question again gives the same cut. The paper is asked once more for a condensed answer,
    # and the extraction says so: fewer rows were kept than the paper reports.
    condensed = False
    try:
        content = _json_content(_ask(system), f"extraction {row.get('id')}")
    except LLMTruncated:
        system = system + _CONDENSED_RULE
        condensed = True
        content = _json_content(_ask(system), f"extraction {row.get('id')} (condensed)")
    parsed = parse_extraction(json.loads(content), source_text)
    # The record of how this was made: the model, the prompt, and when. An extraction that
    # cannot say which model and prompt produced it cannot be reproduced or compared.
    parsed.update({"v": EXTRACTION_VERSION, "source": source, "truncated": truncated,
                   "condensed": condensed,
                   "n_chars": len(source_text), "model": model, "max_tokens": ceiling[0],
                   "prompt_sha": hashlib.sha256(system.encode("utf-8")).hexdigest()[:10],
                   "extracted_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})
    # How much of the paper that text is, and under which limit: a paper cut under a smaller
    # limit than the current one is read again (`_needs_extraction`).
    for k in ("chars_read", "chars_total", "max_chars", "references_dropped"):
        if text_info and k in text_info:
            parsed[k] = text_info[k]
    return parsed


def _cut_under_smaller_limit(ex: dict) -> bool:
    """The text this extraction read was cut, at a limit below today's: more of the paper
    can now be read. An extraction made before the limit was recorded was cut at 60,000."""
    if ex.get("truncated") is not True:
        return False
    try:
        cut_at = int(ex.get("max_chars") or _LEGACY_MAX_CHARS)
    except (TypeError, ValueError):
        cut_at = _LEGACY_MAX_CHARS
    return cut_at < EXTRACTION_MAX_CHARS


def _needs_extraction(row: dict) -> bool:
    """An article is (re)extracted when it has none, when the extraction version moved,
    when the full text arrived after an abstract-only pass, or when its text was cut under a
    smaller limit than the current one. It is skipped once it has failed `_MAX_ATTEMPTS`
    times, and when it has nothing to read.

    "The full text arrived" is asked of the stored CHUNKS (`has_fulltext_chunk`), never of
    the `has_fulltext` flag. That flag is set by the enrichment pass as soon as Unpaywall
    returns an open access LINK, and no text is stored with it. An article flagged that
    way with no full text chunk was therefore selected on every run, read from its
    abstract, written back with `source: "abstract"` and `extraction_attempts` reset to
    zero, then selected again on exactly the same grounds: an unbounded loop, paid every
    time, with 60 000 characters of model budget per article."""
    if int(row.get("extraction_attempts") or 0) >= _MAX_ATTEMPTS:
        return False
    has_chunk = bool(row.get("has_fulltext_chunk"))
    has_text = has_chunk or len((row.get("abstract") or "").strip()) >= _MIN_ABSTRACT_CHARS
    if not has_text:
        return False
    ex = row.get("extraction_json")
    if not isinstance(ex, dict):
        return True
    if int(ex.get("v") or 0) < EXTRACTION_VERSION:
        return True
    if has_chunk and _cut_under_smaller_limit(ex):
        return True
    return ex.get("source") == "abstract" and has_chunk


_ARTICLE_ROWS_SQL = f"""
    SELECT d.id, d.title, d.abstract, d.doi, d.year, d.authors, d.has_fulltext, d.country,
           d.extraction_json, COALESCE(d.extraction_attempts, 0) AS extraction_attempts,
           -- « On DÉTIENT le texte », qui n'est pas « un lien existe » : c'est la
           -- présence d'un morceau de texte intégral, la même question que se pose
           -- scenario_store.scenario_counts_sql.
           EXISTS (SELECT 1 FROM document_chunk c
                    WHERE c.document_id = d.id
                      AND c.chunk_type IN ('fulltext_section', 'full_text'))
               AS has_fulltext_chunk,
           {screening_status_sql('d', 'ars')} AS screening_status
    FROM literature_document d
    JOIN article_scenarios ars ON ars.document_id = d.id
    WHERE ars.scenario_id = :sid AND {{gate}}
    ORDER BY d.id
"""


def _relevant_rows(scenario_id: str) -> list[dict]:
    """EVERY relevant article of the scenario (above the threshold, or included by hand,
    never the excluded): the same gate as the rest of the application."""
    thr = _get_scenario_threshold(scenario_id)
    with engine.connect() as conn:
        return [dict(r) for r in conn.execute(
            text(_ARTICLE_ROWS_SQL.format(gate=relevant_gate_sql("d", "ars", ":thr"))),
            {"sid": scenario_id, "thr": thr}).mappings().all()]


# ─────────────────────────────────────────────────────────────────────────────
# The background job
# ─────────────────────────────────────────────────────────────────────────────
_jobs_lock = threading.Lock()
_jobs: dict[str, dict[str, Any]] = {}


def _llm_client():
    """The LLM client, in one place that the tests replace."""
    from llm_usage import MeteredOpenAI as _OAI
    return _OAI(timeout=_LLM_TIMEOUT_S)


#: The per-job counters: `transient` failures cost no attempt, `skipped` articles were not
#: tried because the job had paused (`paused` says why: the provider looked down).
_JOB_COUNTS = ("total", "done", "failed", "transient", "skipped")


def _extract_scenario(scenario_id: str, disease_hint: str | None, max_articles: int = 0) -> dict[str, int]:
    rows = [r for r in _relevant_rows(scenario_id) if _needs_extraction(r)]
    cap = int(max_articles or EXTRACTION_MAX_ARTICLES or 0)
    if cap > 0:
        rows = rows[:cap]
    with _jobs_lock:
        _jobs[scenario_id].update({**{k: 0 for k in _JOB_COUNTS}, "total": len(rows), "paused": None})
    if not rows:
        return {"total": 0, "done": 0, "failed": 0}
    client = _llm_client()
    # The scenario's own codebook steers how the model names its groups.
    system = _EXTRACTION_SYSTEM + vocabulary_prompt(get_codebook(scenario_id)["nodes"])
    # An outage is not the article's fault. A refused call (the switch, the budget), a quota,
    # a timeout or a 5xx used to count as one of the article's three attempts, so one bad
    # night gave up the whole queue; now they cost nothing, and when the provider looks down
    # the job stops instead of failing through every article left.
    watch = ProviderWatch(cooldown=_openai_in_cooldown)

    def _work(row: dict) -> str:
        if watch.stopped():
            return "skipped"
        try:
            with engine.connect() as conn:
                info = _article_text_info(conn, row)
            result = extract_article(client, row, info["text"], info["source"], info["truncated"],
                                     disease_hint, system, text_info=info)
            watch.success()
            with engine.begin() as conn:
                conn.execute(text("""
                    UPDATE literature_document
                    SET extraction_json = CAST(:j AS jsonb), extraction_at = now(),
                        extraction_attempts = 0
                    WHERE id = :id
                """), {"j": json.dumps(result, ensure_ascii=False), "id": row["id"]})
            return "done"
        except Exception as e:                               # noqa: BLE001 - one article never stops the run
            if is_transient_llm_error(e):
                if _is_openai_quota_error(e):
                    _trip_openai_cooldown()
                watch.failure(e)
                logger.warning(f"extraction article {row.get('id')}: {e} (transient: no attempt charged)")
                return "transient"
            watch.success()                                  # the provider answered; the article failed
            logger.warning(f"extraction article {row.get('id')}: {e}")
            try:
                with engine.begin() as conn:
                    conn.execute(text("UPDATE literature_document SET extraction_attempts = "
                                      "COALESCE(extraction_attempts, 0) + 1 WHERE id = :id"),
                                 {"id": row["id"]})
            except Exception as e2:                          # noqa: BLE001
                logger.warning(f"extraction attempts {row.get('id')}: {e2}")
            return "failed"

    with ThreadPoolExecutor(max_workers=_EXTRACTION_WORKERS) as ex:
        for outcome in ex.map(_work, rows):
            with _jobs_lock:
                _jobs[scenario_id][outcome] += 1
    with _jobs_lock:
        if watch.reason:
            _jobs[scenario_id]["paused"] = watch.reason
            logger.warning(f"extraction {scenario_id} paused, provider unavailable: {watch.reason}")
        return {k: _jobs[scenario_id][k] for k in ("total", "done", "failed")}


def start_extraction(scenario_id: str, disease_hint: str | None = None, max_articles: int = 0) -> dict[str, Any]:
    """Start the extraction in the background, one run per scenario at a time. Nothing is
    started while no call could be made (the master switch, the day's budget, the quota
    circuit breaker): `paused`, with the reason."""
    if not os.getenv("OPENAI_API_KEY"):
        return {"status": "no_llm"}
    probe = ProviderWatch(cooldown=_openai_in_cooldown)
    if probe.preflight("extraction"):
        return {"status": "paused", "reason": probe.reason}
    with _jobs_lock:
        if _jobs.get(scenario_id, {}).get("running"):
            return {"status": "running", **{k: _jobs[scenario_id].get(k, 0) for k in ("total", "done", "failed")}}
        _jobs[scenario_id] = {"running": True, "started_at": time.time(), "total": 0, "done": 0, "failed": 0}

    def _run():
        try:
            _extract_scenario(scenario_id, disease_hint, max_articles)
        except Exception as e:                               # noqa: BLE001
            logger.warning(f"extraction {scenario_id}: {e}")
            with _jobs_lock:
                _jobs[scenario_id]["error"] = str(e)[:300]
        finally:
            with _jobs_lock:
                _jobs[scenario_id]["running"] = False

    threading.Thread(target=_run, daemon=True, name=f"extraction-{scenario_id}").start()
    return {"status": "started"}


def reset_given_up(scenario_id: str) -> int:
    """Give back their attempts to the scenario's relevant articles that failed three times
    and hold no extraction. Returns how many."""
    ids = [int(r["id"]) for r in _relevant_rows(scenario_id)
           if int(r.get("extraction_attempts") or 0) >= _MAX_ATTEMPTS
           and not isinstance(r.get("extraction_json"), dict)]
    if not ids:
        return 0
    with engine.begin() as conn:
        conn.execute(text("UPDATE literature_document SET extraction_attempts = 0 "
                          "WHERE id = ANY(CAST(:ids AS bigint[]))"), {"ids": ids})
    return len(ids)


def extraction_status(scenario_id: str) -> dict[str, Any]:
    """Where the scenario stands, counted over ALL its relevant articles."""
    rows = _relevant_rows(scenario_id)
    done = [r for r in rows if isinstance(r.get("extraction_json"), dict)]
    pending = [r for r in rows if _needs_extraction(r)]
    with _jobs_lock:
        job = dict(_jobs.get(scenario_id) or {})
    return {
        "scenario_id": scenario_id,
        "n_relevant": len(rows),
        "n_extracted": len(done),
        "n_from_fulltext": sum(1 for r in done if r["extraction_json"].get("source") == "fulltext"),
        "n_from_abstract": sum(1 for r in done if r["extraction_json"].get("source") == "abstract"),
        "n_pending": len(pending),
        # Of the pending, those already extracted that the next run reads again: their full
        # text arrived, or their text was cut under a smaller limit than today's.
        "n_to_reread": sum(1 for r in pending if isinstance(r.get("extraction_json"), dict)),
        "n_observations": sum(len(r["extraction_json"].get("observations") or []) for r in done),
        "n_given_up": sum(1 for r in rows if int(r.get("extraction_attempts") or 0) >= _MAX_ATTEMPTS
                          and not isinstance(r.get("extraction_json"), dict)),
        # Read through the condensed retry: at most _CONDENSED_MAX_OBSERVATIONS rows kept.
        "n_condensed": sum(1 for r in done if r["extraction_json"].get("condensed") is True),
        # Read from a text cut at the size limit: the end of the paper was not read.
        "n_truncated": sum(1 for r in done if r["extraction_json"].get("truncated") is True),
        "max_chars": EXTRACTION_MAX_CHARS,
        # Rows of the answers that named no sheet the template has, or no covariate.
        "n_rows_dropped": sum(int(r["extraction_json"].get("n_rows_dropped") or 0) for r in done),
        "running": bool(job.get("running")),
        "job": ({**{k: job.get(k, 0) for k in _JOB_COUNTS}, "paused": job.get("paused")}
                if job else None),
        "version": EXTRACTION_VERSION,
    }


# ─────────────────────────────────────────────────────────────────────────────
# The REDUCE half: counts over EVERY relevant article, in SQL, no LLM
# ─────────────────────────────────────────────────────────────────────────────
# What the assistant and the reviewers need is not a paragraph written from 24 excerpts
# but "how many of the relevant papers report sex-disaggregated data", counted over all of
# them. The map step above wrote the per-article facts; this aggregates them, the way
# api/digest.py does for the concepts.
#
# Two things the figures must never hide. The DENOMINATOR: only articles already extracted
# can be counted, so every count says out of how many, and how many relevant articles are
# still unread. And the SOURCE: a paper extracted from its abstract alone cannot show a
# table, so its "false" is weaker than a "false" read from the full text. Both are reported
# next to the count, not behind it.

_COVERAGE_LABELS_FR = {
    "sex_gender": "sexe ou genre", "age": "age", "occupation": "profession",
    "kap_risk_perception": "connaissances, attitudes, pratiques ou perception du risque",
    "ppe": "equipements de protection", "vaccination": "vaccination",
    "human_testing": "tests chez l'humain", "animal_host": "hotes animaux",
    "environment": "environnement", "vector": "vecteurs",
}
_TOP_GROUPS = 12


def normalise_group(index, sheet: str, group: str) -> str:
    """A group as the codebook names it (its level-1 key), else the folded text."""
    return normalise(index, sheet or "", group, "")["group_key"] or (group or "").lower()
_TOP_CRUDE = 12

_OBS_FROM = """
    FROM literature_document d
    JOIN article_scenarios ars ON ars.document_id = d.id
    CROSS JOIN LATERAL jsonb_array_elements(d.extraction_json->'observations') AS o
    WHERE ars.scenario_id = :sid AND {gate}
      AND jsonb_typeof(d.extraction_json->'observations') = 'array'
"""


def extraction_digest(scenario_id: str, threshold: float | None = None) -> dict[str, Any]:
    """Counts over ALL the relevant articles of what the extraction found.

    `n_extracted` is the denominator of every proportion below it; `n_relevant` -
    `n_extracted` articles have not been read yet and say nothing. `coverage[key]` counts
    the extracted articles that report the item, and `coverage_fulltext[key]` the ones read
    from the full text. `crude_counts` sums cases and population over rows that report both
    and whose quote was found in the text: it is a tally, NOT a pooled estimate (the
    populations differ, the labels are free text until the codebook, and one population can
    appear in two rows), and the digest says so."""
    thr = _get_scenario_threshold(scenario_id) if threshold is None else float(threshold)
    gate = relevant_gate_sql("d", "ars", ":thr")
    out: dict[str, Any] = {"scenario_id": scenario_id, "threshold": thr, "version": EXTRACTION_VERSION}
    try:
        cov_sql = ",\n".join(
            f"COUNT(*) FILTER (WHERE d.extraction_json->'coverage'->>'{k}' = 'true') AS c_{k}, "
            f"COUNT(*) FILTER (WHERE d.extraction_json->'coverage'->>'{k}' = 'true' "
            f"AND d.extraction_json->>'source' = 'fulltext') AS f_{k}" for k in COVERAGE_KEYS)
        with engine.connect() as conn:
            head = _rows(conn, f"""
                SELECT COUNT(*) AS n_relevant,
                       COUNT(*) FILTER (WHERE jsonb_typeof(d.extraction_json) = 'object') AS n_extracted,
                       COUNT(*) FILTER (WHERE d.extraction_json->>'source' = 'fulltext') AS n_fulltext,
                       COUNT(*) FILTER (WHERE d.extraction_json->>'source' = 'abstract') AS n_abstract,
                       COUNT(*) FILTER (WHERE d.extraction_json->>'truncated' = 'true') AS n_truncated,
                       COUNT(*) FILTER (WHERE d.extraction_json->>'condensed' = 'true') AS n_condensed,
                       {cov_sql}
                FROM literature_document d
                JOIN article_scenarios ars ON ars.document_id = d.id
                WHERE ars.scenario_id = :sid AND {gate}
            """, scenario_id, thr)[0]
            out["n_relevant"] = int(head["n_relevant"] or 0)
            out["n_extracted"] = int(head["n_extracted"] or 0)
            out["n_fulltext"] = int(head["n_fulltext"] or 0)
            out["n_abstract"] = int(head["n_abstract"] or 0)
            # Read from a text cut at the size limit (its end unread), and through the
            # condensed retry (at most 50 rows kept): fewer rows than the paper reports.
            out["n_truncated"] = int(head["n_truncated"] or 0)
            out["n_condensed"] = int(head["n_condensed"] or 0)
            out["max_chars"] = EXTRACTION_MAX_CHARS
            out["n_unread"] = out["n_relevant"] - out["n_extracted"]
            out["coverage"] = {k: int(head[f"c_{k}"] or 0) for k in COVERAGE_KEYS}
            out["coverage_fulltext"] = {k: int(head[f"f_{k}"] or 0) for k in COVERAGE_KEYS}

            out["by_model"] = [
                {"model": r["model"], "prompt_sha": r["prompt_sha"], "n": int(r["n"])}
                for r in _rows(conn, f"""
                    SELECT COALESCE(d.extraction_json->>'model', 'unknown') AS model,
                           COALESCE(d.extraction_json->>'prompt_sha', 'unknown') AS prompt_sha,
                           COUNT(*) AS n
                    FROM literature_document d JOIN article_scenarios ars ON ars.document_id = d.id
                    WHERE ars.scenario_id = :sid AND {gate} AND jsonb_typeof(d.extraction_json) = 'object'
                    GROUP BY 1, 2 ORDER BY n DESC
                """, scenario_id, thr)]
            out["by_sheet"] = [
                {"sheet": r["sheet"], "n_rows": int(r["n_rows"]), "n_articles": int(r["n_articles"]),
                 "n_quote_found": int(r["n_verified"])}
                for r in _rows(conn, f"""
                    SELECT o->>'sheet' AS sheet, COUNT(*) AS n_rows, COUNT(DISTINCT d.id) AS n_articles,
                           COUNT(*) FILTER (WHERE o->>'quote_verified' = 'true') AS n_verified
                    {_OBS_FROM.format(gate=gate)}
                    GROUP BY 1 ORDER BY n_rows DESC
                """, scenario_id, thr)]

            index = get_index(scenario_id)
            ids_by_group: dict[tuple[str, str], set] = {}
            for r in _rows(conn, f"""
                SELECT o->>'sheet' AS sheet, o->>'group' AS grp, array_agg(DISTINCT d.id) AS ids
                {_OBS_FROM.format(gate=gate)} AND NULLIF(TRIM(o->>'group'), '') IS NOT NULL
                GROUP BY 1, 2
            """, scenario_id, thr):
                key = normalise_group(index, r["sheet"], r["grp"])
                ids_by_group.setdefault((r["sheet"], key), set()).update(r["ids"])
            out["top_groups"] = [
                {"sheet": sh, "value": v, "n": len(ids)}
                for (sh, v), ids in sorted(ids_by_group.items(), key=lambda kv: (-len(kv[1]), kv[0]))[:_TOP_GROUPS]]

            crude = _rows(conn, f"""
                SELECT o->>'sheet' AS sheet, LOWER(TRIM(o->>'covariate')) AS covariate,
                       COUNT(DISTINCT d.id) AS n_studies,
                       SUM((o->>'n_cases')::numeric) AS n_cases,
                       SUM((o->>'pop_risk')::numeric) AS pop_risk
                {_OBS_FROM.format(gate=gate)}
                  AND jsonb_typeof(o->'n_cases') = 'number' AND jsonb_typeof(o->'pop_risk') = 'number'
                  AND o->>'quote_verified' = 'true'
                GROUP BY 1, 2 ORDER BY n_studies DESC, pop_risk DESC, 2 LIMIT :top
            """, scenario_id, thr, top=_TOP_CRUDE)
            out["crude_counts"] = [
                {"sheet": r["sheet"], "covariate": r["covariate"], "n_studies": int(r["n_studies"]),
                 "n_cases": float(r["n_cases"] or 0), "pop_risk": float(r["pop_risk"] or 0)}
                for r in crude]
        out["crude_counts_note"] = ("A tally of rows that report both cases and population with a quote "
                                    "found in the text. Not a pooled estimate: populations differ, labels "
                                    "are free text, and one population may appear in two rows.")
        out["complete"] = True
    except Exception as e:                                   # noqa: BLE001 - never blocks the caller
        logger.warning(f"extraction_digest {scenario_id}: {e}")
        out.update({"complete": False, "error": str(e)[:300]})
    return out


def extraction_to_prompt(d: dict | None, max_chars: int = 6000) -> str:
    """The extraction digest as a block for a prompt, or "" when there is nothing to say
    (nothing extracted, or the aggregation failed: an incomplete block must assert nothing).

    It states its own denominator and its source, and tells the model to quote them, so
    "34 papers report sex-disaggregated data" is never read as "34 of all the relevant".

    The instructions come first and the block is cut by whole sections, never mid-line: at
    2,200 characters it was cut from the end, and the crude counts went first, then the
    last line, the very instruction to quote the denominator."""
    from .digest import fit_prompt_sections                  # the same whole-section cut as the digest
    if not d or not d.get("complete") or not d.get("n_extracted"):
        return ""
    n, rel = d["n_extracted"], d["n_relevant"]
    sections = [
        ("Extraction", f"EXTRACTION STRUCTUREE: {n} des {rel} articles pertinents ont ete lus ({d['n_fulltext']} "
                       f"depuis le texte integral, {d['n_abstract']} depuis le resume seul, qui ne contient pas les "
                       f"tableaux : pour eux un element absent est un minorant, pas une absence)."),
        ("Denominateur", f"Toute proportion tiree de ce bloc doit citer son denominateur ({n} articles "
                         f"extraits sur {rel})."),
    ]
    if d.get("n_unread"):
        sections.append(("Non extraits", f"{d['n_unread']} articles pertinents ne sont pas encore extraits : "
                                         f"ne rien conclure sur eux."))
    mandatory = len(sections)
    cov = [f"{_COVERAGE_LABELS_FR[k]} {d['coverage'][k]} (dont {d['coverage_fulltext'][k]} en texte integral)"
           for k in COVERAGE_KEYS if d["coverage"].get(k)]
    sections.append(("Couverture", f"Parmi les {n} articles extraits, rapportent des donnees sur : "
                     + (", ".join(cov) if cov else "aucun des elements suivis (sexe, age, profession, CAP, EPI, vaccination...)") + "."))
    if d.get("by_sheet"):
        sections.append(("Lignes extraites", "Lignes extraites : " + ", ".join(
            f"{r['sheet']} {r['n_rows']} ({r['n_articles']} articles, citation retrouvee {r['n_quote_found']})"
            for r in d["by_sheet"]) + "."))
    if d.get("top_groups"):
        sections.append(("Groupes", "Groupes les plus frequents : " + ", ".join(
            f"{r['value']} [{r['sheet']}] ({r['n']})" for r in d["top_groups"]) + "."))
    if d.get("crude_counts"):
        sections.append(("Totaux bruts", "Totaux bruts (lignes avec cas et population, citation retrouvee ; pas une estimation poolee) : "
                         + "; ".join(f"{r['covariate']} [{r['sheet']}] {r['n_studies']} etudes, "
                                     f"{r['n_cases']:g} cas sur {r['pop_risk']:g}" for r in d["crude_counts"]) + "."))
    return fit_prompt_sections(sections, max_chars, mandatory=mandatory)


# ─────────────────────────────────────────────────────────────────────────────
# Export in the template's shape
# ─────────────────────────────────────────────────────────────────────────────
def _first_author(authors: str | None) -> str:
    """The first author's family name when the stored form makes it clear ("Dressler A",
    "Dressler, Aparna"); otherwise the name as stored, for the reviewer to fix."""
    first = re.split(r";|\n", (authors or "").strip())[0].strip()
    if "," in first:
        head = first.split(",")[0].strip()
        if head and len(head.split()) <= 2:
            return head
    toks = first.split()
    if len(toks) >= 2 and re.fullmatch(r"[A-Z]{1,3}", toks[-1]):
        return " ".join(toks[:-1])
    return first


def _cell(v: Any) -> Any:
    if isinstance(v, bool):
        return "Y" if v else "N"
    if isinstance(v, str):
        # A control character from a PDF text in one quote made openpyxl refuse the whole file.
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def template_rows(articles: list[dict], id_prefix: str = "LR", index=None, reviews: dict | None = None,
                  include_rejected: bool = False, nuts_index=None) -> dict[str, list[dict]]:
    """The extractions as template rows: `{"ref": [...], "human_susc": [...], ...}`.
    `articles` are relevant-article rows carrying an `extraction_json`. Pure."""
    out: dict[str, list[dict]] = {"ref": [], **{s: [] for s in SHEETS}}
    for a in articles:
        ex = a.get("extraction_json")
        if not isinstance(ex, dict):
            continue
        aid = f"{id_prefix}{a['id']}"
        ref = ex.get("ref") or {}
        doi = (a.get("doi") or "").strip()
        place = resolve_location(ref.get("location"), a.get("country"), nuts_index) if nuts_index is not None else {}
        out["ref"].append({
            "id": aid, "first_author": _first_author(a.get("authors")),
            "reference": f"https://doi.org/{doi}" if doi and not doi.startswith("http") else doi,
            "year": a.get("year"), "description": ref.get("description"),
            "article_type": ref.get("article_type"), "study_start": ref.get("study_start"),
            "study_end": ref.get("study_end"), "location": ref.get("location"),
            "nuts1": place.get("nuts1"), "nuts2": place.get("nuts2"), "nuts3": place.get("nuts3"),
            "notes_geo": ref.get("notes_geo"),
            "risk_pop": ref.get("risk_pop"), "positive": ref.get("positive"),
            "percent_positive": ref.get("percent_positive"),
            "math_model": ref.get("math_model"), "model_type": ref.get("model_type"),
            "exclusion": (a.get("screening_reason") if a.get("screening_status") == "excluded" else None),
            "source": ex.get("source"), "chars_read": ex.get("chars_read", ex.get("n_chars")),
            "text_cut": ex.get("truncated") is True, "condensed": ex.get("condensed") is True,
        })
        observations = [o for o in (ex.get("observations") or []) if isinstance(o, dict)]
        if reviews is not None:
            reviewed, _stale = overlay(int(a["id"]), observations, reviews)
        else:
            reviewed = [{**o, "review_status": "unreviewed", "reviews": [], "effective": o} for o in observations]
        for o in reviewed:
            if o["review_status"] == "rejected" and not include_rejected:
                continue
            fields = o["effective"]                              # the reviewers' corrections applied
            sheet = fields.get("sheet")
            if sheet in out:
                row = annotate(index, fields) if index is not None else dict(fields)
                out[sheet].append({**row, "id": aid, "source": ex.get("source"),
                                   "text_truncated": ex.get("truncated") is True,
                                   "condensed": ex.get("condensed") is True,
                                   "review_status": o["review_status"],
                                   "reviewed_by": ", ".join(r["reviewer"] for r in o["reviews"]) or None})
    return out


def _models_line(articles: list[dict]) -> str:
    """'model, prompt abc123 (12 papers); ...' from what each extraction says about itself."""
    c: Counter = Counter()
    for a in articles:
        ex = a.get("extraction_json")
        if isinstance(ex, dict):
            c[(ex.get("model") or "unknown", ex.get("prompt_sha") or "unknown")] += 1
    return "; ".join(f"{m}, prompt {p} ({n} papers)" for (m, p), n in c.most_common())


def _read_counts(articles: list[dict]) -> tuple[int, int, int]:
    """(cut at the size limit, condensed, rows dropped) over the extracted articles."""
    exs = [a["extraction_json"] for a in articles if isinstance(a.get("extraction_json"), dict)]
    return (sum(1 for ex in exs if ex.get("truncated") is True), sum(1 for ex in exs if ex.get("condensed") is True),
            sum(int(ex.get("n_rows_dropped") or 0) for ex in exs))


#: The geography in full: one row per country, then one per region it holds.
COUNTRY_SHEET_HEADER = ["Country", "ISO 3166-1 alpha-2", "NUTS 0", "Papers", "Rows", "NUTS region", "Region name",
                        "Papers in the region"]


def build_workbook(articles: list[dict], coverage_note: str, id_prefix: str = "LR", index=None,
                   reviews: dict | None = None, include_rejected: bool = False, nuts_index=None,
                   geo: dict | None = None, pooled: dict | None = None) -> bytes:
    """An .xlsx with the template's sheets and column titles, plus review columns on the
    right and a README sheet that says what the numbers cover.

    With `geo` (the scenario's geography) and `pooled` (a `full` pooled result), three more
    sheets hold the lists that the report and the panel cut: every country with its regions
    (COUNTRIES), every label with its studies, pooled or too few (POOLED), and every comparison
    between two labels (COMPARISONS)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = template_rows(articles, id_prefix, index, reviews, include_rejected, nuts_index)
    n_cut, n_condensed, n_dropped = _read_counts(articles)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    lines = [
        "LiteRev structured extraction, shaped like the T2.4 / T4.5 data extraction template.",
        coverage_note,
        "Every row was produced by a model and must be checked against the paper by a reviewer.",
        "'Quote found in text' = N means the model's quote was not found in the text it read: "
        "check that row first.",
        "'Extracted from' = abstract means the paper's full text was not available: tables "
        "are not in abstracts, so most values are missing, not absent from the paper.",
        f"{n_cut} papers were read from a text cut at the size limit ({EXTRACTION_MAX_CHARS:,} characters "
        "per paper today): the end of their text, often tables and supplementary material, was not "
        "read. REF says which ('Text cut at the size limit') and how much was read ('Characters "
        "read'). The reference list of a full text is not read.",
        f"{n_condensed} papers held more data than one answer of the model can carry and were read "
        f"in a condensed form (at most {_CONDENSED_MAX_OBSERVATIONS} rows each, rows with counts "
        "first): some of their rows are missing. REF says which ('Condensed answer').",
        "NUTS columns are filled from the place name where it matches the NUTS list in use: the "
        "country always, a region only when the list knows it (load Eurostat's NUTS file for the "
        "whole of Europe). A blank means the place could not be resolved, not that it is nowhere.",
        "'Review status' says what the reviewers decided (accepted, edited, rejected, conflict or "
        "unreviewed). Edited values are the reviewers' corrections; rejected rows are left out "
        "unless the file was exported with them.",
        "ID is LR + the LiteRev article id; rename it to your initials and a number if needed.",
        "Made with: " + (_models_line(articles) or "unknown") + ".",
    ]
    if n_dropped:
        lines.insert(7, f"{n_dropped} rows of the model's answers named no sheet of the template, or no "
                        "covariate, and are not in this file.")
    if geo is not None:
        lines.append(f"COUNTRIES: every country of the {geo.get('n_papers', 0)} extracted papers, with its "
                     f"regions; {geo.get('n_unresolved', 0)} papers name no place that resolves to one country.")
    if pooled is not None:
        f = pooled.get("filters") or {}
        lines.append(f"POOLED: every codebook label with its studies ({len(pooled['pooled'])}), pooled when it has "
                     f"at least {f.get('min_studies', 3)} studies. COMPARISONS: every comparison between two "
                     f"labels of a group ({len(pooled['comparisons'])}). Rows whose quote was not found are "
                     "left out" + (", and diseases are kept apart." if f.get("split_disease", True) else "."))
    for line in lines:
        readme.append([line])
    readme.column_dimensions["A"].width = 120
    bold = Font(bold=True)

    ws = wb.create_sheet("REF")
    ws.append([c[0] for c in REF_COLUMNS] + [c[0] for c in REF_READ_COLUMNS])
    for r in rows["ref"]:
        ws.append([_cell(r.get(f)) for _, f in REF_COLUMNS] + [_cell(r.get(f)) for _, f in REF_READ_COLUMNS])
    for key in SHEETS:
        title, cols = SHEET_COLUMNS[key]
        have = {f for _, f in cols}
        extra = [(h, f) for h, f in REVIEW_COLUMNS if f not in have]
        ws = wb.create_sheet(title)
        ws.append([c[0] for c in cols] + [c[0] for c in extra])
        for r in rows[key]:
            ws.append([_cell(r.get(f)) for _, f in cols] + [_cell(r.get(f)) for _, f in extra])
    if geo is not None:
        ws = wb.create_sheet("COUNTRIES")
        ws.append(COUNTRY_SHEET_HEADER)
        for c in geo.get("countries") or []:
            ws.append([_cell(c.get("name")), c.get("iso2"), c.get("nuts0"), c.get("n_papers"), c.get("n_rows")])
            for g in c.get("regions") or []:
                ws.append([_cell(c.get("name")), c.get("iso2"), c.get("nuts0"), None, None, g.get("code"),
                           _cell(g.get("name")), g.get("n_papers")])
    if pooled is not None:
        from .pooling import COMPARISON_HEADER, POOLED_LABEL_HEADER, comparison_rows, pooled_label_rows
        for title, header, body in (("POOLED", POOLED_LABEL_HEADER, pooled_label_rows(pooled["pooled"])),
                                    ("COMPARISONS", COMPARISON_HEADER, comparison_rows(pooled["comparisons"]))):
            ws = wb.create_sheet(title)
            ws.append(header)
            for r in body:
                ws.append([_cell(v) for v in r])
    for ws in wb.worksheets[1:]:
        for c in ws[1]:
            c.font = bold
        ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_long_csv(articles: list[dict], id_prefix: str = "LR", index=None, reviews: dict | None = None,
                   include_rejected: bool = False, nuts_index=None) -> str:
    """One flat table, one row per observation, with its sheet: for scripts and datasets.
    `text_truncated` and `condensed` say, on each row, that its paper was read cut or condensed."""
    rows = template_rows(articles, id_prefix, index, reviews, include_rejected, nuts_index)
    fields = ["id", "sheet", "transmission_mode", "disease", "group", "covariate", "l1", "l2", "value", "descr",
              "notes", "n_cases", "pop_risk", "original_name", "page_section", "source_kind",
              "quote", "quote_verified", "review_status", "reviewed_by", "source", "text_truncated", "condensed"]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(fields)
    for key in SHEETS:
        for r in rows[key]:
            w.writerow([key if f == "sheet" else r.get(f) for f in fields])
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Endpoints
# ─────────────────────────────────────────────────────────────────────────────
@app.post("/user-scenarios/{scenario_id}/extraction/run")
def run_scenario_extraction(scenario_id: str, max_articles: int = 0, retry_given_up: bool = False,
                            _: None = Depends(require_api_key)) -> dict[str, Any]:
    """Extract, in the background, every relevant article that has no extraction yet (or
    whose full text arrived since an abstract-only pass, or whose text was cut under a smaller
    limit than today's). No cap unless `max_articles` or EXTRACTION_MAX_ARTICLES says so.
    `status` is `paused`, with a `reason`, when no call can be made (the master switch, the
    day's budget, the quota circuit breaker).

    `retry_given_up` first gives the articles left alone after three failures their
    attempts back, for when the cause of the failures has been fixed (the answer ceiling
    that cut the richest papers, for one). Only this scenario's relevant articles."""
    scenario = _get_user_scenario_or_404(scenario_id)
    n_reset = reset_given_up(scenario_id) if retry_given_up else 0
    out = start_extraction(scenario_id, scenario.get("name"), max(0, int(max_articles or 0)))
    if retry_given_up:
        out = {**out, "n_reset": n_reset}
    return {**out, **{k: v for k, v in extraction_status(scenario_id).items()
                      if k in ("n_relevant", "n_pending")}}


@app.get("/user-scenarios/{scenario_id}/extraction/status")
def get_scenario_extraction_status(scenario_id: str) -> dict[str, Any]:
    _get_user_scenario_or_404(scenario_id)
    return extraction_status(scenario_id)


def reverify_quotes(scenario_id: str) -> dict[str, Any]:
    """Check the quote of every stored row of the scenario's relevant extractions again, against
    the paper's WHOLE text (title, abstract, full text, never cut) and with today's check, which
    undoes what a PDF does to text. No model call: only `quote_verified` changes (and the date
    of the check is recorded), so the reviewers' decisions keep their rows. Counts before and
    after."""
    rows = [r for r in _relevant_rows(scenario_id) if isinstance(r.get("extraction_json"), dict)]
    out = {"scenario_id": scenario_id, "n_articles": len(rows), "n_articles_changed": 0, "n_observations": 0,
           "verified_before": 0, "verified_after": 0, "newly_found": 0, "no_longer_found": 0}
    for r in rows:
        ex = r["extraction_json"]
        obs = ex.get("observations") if isinstance(ex.get("observations"), list) else []
        if not obs:
            continue
        with engine.connect() as conn:
            whole = _article_whole_text(conn, r)
        changed = False
        for o in obs:
            if not isinstance(o, dict):
                continue
            before, after = o.get("quote_verified") is True, quote_in_text(o.get("quote"), whole)
            out["n_observations"] += 1
            out["verified_before"] += before
            out["verified_after"] += after
            if before != after:
                out["newly_found" if after else "no_longer_found"] += 1
                o["quote_verified"] = after
                changed = True
        if changed:
            ex["quotes_rechecked_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
            with engine.begin() as conn:
                # Only if the extraction is still the one that was read: a run may have replaced it.
                n = conn.execute(text("""
                    UPDATE literature_document SET extraction_json = CAST(:j AS jsonb)
                    WHERE id = :id AND extraction_json->>'extracted_at' IS NOT DISTINCT FROM :at
                """), {"j": json.dumps(ex, ensure_ascii=False), "id": r["id"],
                       "at": ex.get("extracted_at")}).rowcount
            out["n_articles_changed"] += int(bool(n))
    return out


@app.post("/user-scenarios/{scenario_id}/extraction/reverify")
def reverify_scenario_quotes(scenario_id: str, _: None = Depends(require_api_key)) -> dict[str, Any]:
    """Re-check the quotes of the scenario's stored extractions (see `reverify_quotes`), with no
    model call. 409 while an extraction run is writing them."""
    _get_user_scenario_or_404(scenario_id)
    with _jobs_lock:
        if _jobs.get(scenario_id, {}).get("running"):
            raise HTTPException(status_code=409, detail="An extraction run is in progress: re-check the quotes when it ends.")
    return reverify_quotes(scenario_id)


# ─────────────────────────────────────────────────────────────────────────────
# Seeing a quote in the paper
# ─────────────────────────────────────────────────────────────────────────────
_WORD = re.compile(r"[^\W_]+", re.UNICODE)
_MIN_PARTIAL_WORDS = 4


def locate_quote(source_text: str, quote: str | None) -> dict[str, Any] | None:
    """Where a quote sits in the text: `{start, end, partial}` or None.

    Words and figures must match in order, and anything that is not a letter or a digit may
    stand between them, which is what line breaks, hyphenation and spacing in a PDF text look
    like. When the whole quote is not there, the longest run of its first or last words that is
    (at least four) is returned as `partial`, so a reviewer can see where the model most likely
    took it from and how it differs. The first occurrence is returned."""
    words = _WORD.findall(quote or "")
    if not words or not source_text:
        return None

    def find(ws: list[str]) -> tuple[int, int] | None:
        m = re.search(r"\W+".join(re.escape(w) for w in ws), source_text, re.IGNORECASE)
        return (m.start(), m.end()) if m else None

    whole = find(words)
    if whole:
        return {"start": whole[0], "end": whole[1], "partial": False}
    for n in range(len(words) - 1, _MIN_PARTIAL_WORDS - 1, -1):
        for part in (words[:n], words[-n:]):
            hit = find(part)
            if hit:
                return {"start": hit[0], "end": hit[1], "partial": True}
    return None


@app.get("/user-scenarios/{scenario_id}/articles/{article_id}/text")
def get_article_text_window(scenario_id: str, article_id: int, quote: str | None = Query(None, max_length=600),
                            context: int = Query(700, ge=0, le=6000)) -> dict[str, Any]:
    """The text the model read for this article, around a quote: `before`, `match`, `after`.

    `source` is `fulltext` or `abstract`, as for the extraction. `found` says whether the quote
    is in the text, `partial` whether only part of it is; without a match the start of the text
    is returned. The text is the same one the quote was checked against, cut at the same size."""
    _get_user_scenario_or_404(scenario_id)
    with engine.connect() as conn:
        row = conn.execute(text("""
            SELECT d.id, d.title, d.abstract FROM literature_document d
            JOIN article_scenarios ars ON ars.document_id = d.id
            WHERE ars.scenario_id = :sid AND d.id = :id
        """), {"sid": scenario_id, "id": article_id}).mappings().first()
        if not row:
            raise HTTPException(status_code=404, detail="Article not in this scenario.")
        body, source, truncated = _article_text(conn, dict(row))
    hit = locate_quote(body, quote)
    if hit:
        start, end = hit["start"], hit["end"]
    else:
        start = end = 0
    lo, hi = max(0, start - context), min(len(body), end + context)
    return {"article_id": article_id, "title": row["title"], "source": source, "text_truncated": truncated,
            "n_chars": len(body), "found": bool(hit and not hit["partial"]), "partial": bool(hit and hit["partial"]),
            "start": start, "end": end, "window_start": lo, "window_end": hi,
            "before": body[lo:start], "match": body[start:end], "after": body[end:hi]}


@app.get("/user-scenarios/{scenario_id}/extraction/coverage")
def get_scenario_extraction_coverage(scenario_id: str) -> dict[str, Any]:
    """What the extraction found, counted over ALL the relevant articles (see
    `extraction_digest`): who reports sex, age, KAP, PPE..., and out of how many."""
    _get_user_scenario_or_404(scenario_id)
    return extraction_digest(scenario_id)


def _review_counts(doc_ids: list[int]) -> dict[int, dict[str, int]]:
    """Per article: observations with a decision, rejected, and in conflict. Only the
    articles that have any decision are read."""
    reviews = load_reviews(doc_ids)
    with_reviews = sorted({d for d, _k in reviews})
    if not with_reviews:
        return {}
    from .extraction_review import combine, observation_keys
    with engine.connect() as conn:
        rows = conn.execute(text("SELECT id, extraction_json->'observations' AS obs FROM literature_document "
                                 "WHERE id = ANY(:ids)"), {"ids": with_reviews}).mappings().all()
    out: dict[int, dict[str, int]] = {}
    for r in rows:
        obs = [o for o in (r["obs"] or []) if isinstance(o, dict)] if isinstance(r["obs"], list) else []
        c = {"n_reviewed": 0, "n_rejected": 0, "n_conflict": 0}
        for k in observation_keys(obs):
            st = combine(reviews.get((int(r["id"]), k), []))["status"]
            if st != "unreviewed":
                c["n_reviewed"] += 1
            c["n_rejected"] += st == "rejected"
            c["n_conflict"] += st == "conflict"
        out[int(r["id"])] = c
    return out


#: Same bound as the PICO list: a page is at most this many articles.
EXTRACTION_PAGE_MAX = 5000


@app.get("/user-scenarios/{scenario_id}/extraction/articles")
def list_scenario_extraction_articles(scenario_id: str, limit: int = EXTRACTION_PAGE_MAX,
                                      offset: int = 0) -> dict[str, Any]:
    """EVERY relevant article of the scenario with what its extraction found, paged.

    The extracted ones come first. The rows carry counts, not the observations (read one
    article's with /articles/{id}/extraction), so the page stays small. The response says
    how many articles exist, how many are extracted and whether more pages remain, so the
    screen never reads a page as the whole corpus."""
    _get_user_scenario_or_404(scenario_id)
    limit = max(1, min(int(limit), EXTRACTION_PAGE_MAX))
    offset = max(0, int(offset))
    thr = _get_scenario_threshold(scenario_id)
    gate = relevant_gate_sql("d", "ars", ":thr")
    obs = ("CASE WHEN jsonb_typeof(d.extraction_json->'observations') = 'array' "
           "THEN d.extraction_json->'observations' ELSE '[]'::jsonb END")
    with engine.connect() as conn:
        rows = [dict(r) for r in conn.execute(text(f"""
            SELECT d.id, d.title, d.year, d.doi, d.journal,
                   COALESCE(jsonb_typeof(d.extraction_json) = 'object', FALSE) AS has_extraction,
                   d.extraction_json->>'source' AS source,
                   COALESCE(d.extraction_json->>'truncated', 'false') = 'true' AS text_truncated,
                   COALESCE(d.extraction_json->>'condensed', 'false') = 'true' AS condensed,
                   d.extraction_json->'coverage' AS coverage,
                   jsonb_array_length({obs}) AS n_observations,
                   (SELECT COUNT(*) FROM jsonb_array_elements({obs}) o
                     WHERE o->>'quote_verified' = 'true') AS n_quote_found,
                   COALESCE(d.extraction_attempts, 0) AS attempts
            FROM literature_document d
            JOIN article_scenarios ars ON ars.document_id = d.id
            WHERE ars.scenario_id = :sid AND {gate}
            ORDER BY COALESCE(jsonb_typeof(d.extraction_json) = 'object', FALSE) DESC, d.year DESC NULLS LAST, d.id DESC
            LIMIT :limit OFFSET :offset
        """), {"sid": scenario_id, "thr": thr, "limit": limit, "offset": offset}).mappings().all()]
        head = conn.execute(text(f"""
            SELECT COUNT(*) AS total,
                   COUNT(*) FILTER (WHERE jsonb_typeof(d.extraction_json) = 'object') AS extracted
            FROM literature_document d JOIN article_scenarios ars ON ars.document_id = d.id
            WHERE ars.scenario_id = :sid AND {gate}
        """), {"sid": scenario_id, "thr": thr}).mappings().first()
    counts = _review_counts([int(r["id"]) for r in rows if int(r["n_observations"] or 0) > 0])
    for r in rows:
        r.update(counts.get(int(r["id"]), {"n_reviewed": 0, "n_rejected": 0, "n_conflict": 0}))
        r["n_observations"] = int(r["n_observations"] or 0)
        r["n_quote_found"] = int(r["n_quote_found"] or 0)
        r["coverage"] = {k: bool((r["coverage"] or {}).get(k)) for k in COVERAGE_KEYS} if r["has_extraction"] else None
    total = int(head["total"] or 0)
    returned = len(rows)
    return {
        "scenario_id": scenario_id, "threshold": thr, "total": total,
        "extracted": int(head["extracted"] or 0), "offset": offset, "limit": limit,
        "returned": returned, "truncated": offset + returned < total,
        "next_offset": offset + returned if offset + returned < total else None,
        "page_max": EXTRACTION_PAGE_MAX, "articles": rows,
    }


@app.get("/user-scenarios/{scenario_id}/articles/{article_id}/extraction")
def get_article_extraction(scenario_id: str, article_id: int) -> dict[str, Any]:
    _get_user_scenario_or_404(scenario_id)
    with engine.connect() as conn:
        row = conn.execute(text("""
            SELECT d.id, d.title, d.extraction_json, d.extraction_at
            FROM literature_document d JOIN article_scenarios ars ON ars.document_id = d.id
            WHERE ars.scenario_id = :sid AND d.id = :id
        """), {"sid": scenario_id, "id": article_id}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Article not in this scenario.")
    ex = row["extraction_json"]
    stale: list[dict] = []
    if isinstance(ex, dict) and isinstance(ex.get("observations"), list):
        index = get_index(scenario_id)
        stored = [o for o in ex["observations"] if isinstance(o, dict)]
        obs, stale = overlay_annotated(int(row["id"]), stored, index)
        ex = {**ex, "observations": obs}
    return {"id": row["id"], "title": row["title"], "extracted_at": row["extraction_at"],
            "extraction": ex, "stale_reviews": stale}


@app.get("/user-scenarios/{scenario_id}/extraction/dataset")
def export_extraction_dataset(scenario_id: str, include_unreviewed: bool = Query(False)) -> Response:
    """The annotated dataset, one JSON object per line (JSONL): each observation a reviewer
    accepted or corrected, with the quote it rests on, its codebook label, who validated it, and
    the model and prompt that produced it. This is the labelled data the protocol wants for
    training and evaluating extraction models. Rejected and conflicting rows are never in it;
    unreviewed ones only with `include_unreviewed`."""
    _get_user_scenario_or_404(scenario_id)
    rows = [r for r in _relevant_rows(scenario_id) if isinstance(r.get("extraction_json"), dict)]
    reviews = load_reviews([int(r["id"]) for r in rows])
    index = get_index(scenario_id)
    lines = []
    for a in rows:
        ex = a["extraction_json"]
        stored = [o for o in (ex.get("observations") or []) if isinstance(o, dict)]
        obs, _stale = overlay(int(a["id"]), stored, reviews)
        for o in obs:
            st = o["review_status"]
            if st in ("rejected", "conflict") or (st == "unreviewed" and not include_unreviewed):
                continue
            eff = annotate(index, o["effective"])
            lines.append(json.dumps({
                "article": {"id": a["id"], "title": a.get("title"), "doi": a.get("doi"), "year": a.get("year")},
                "observation": {k: eff.get(k) for k in ("sheet", "group", "covariate", "value", "descr", "n_cases",
                                                        "pop_risk", "disease", "transmission_mode", "page_section",
                                                        "source_kind", "l1", "l2", "label_path")},
                "quote": eff.get("quote"), "quote_found_in_text": eff.get("quote_verified"),
                "review": {"status": st, "reviewers": [r["reviewer"] for r in o["reviews"]]},
                # Whether the paper was read whole: a cut text or a condensed answer holds fewer
                # rows than the paper reports, which a dataset built from this must know.
                "extraction": {"source": ex.get("source"), "model": ex.get("model"),
                               "prompt_sha": ex.get("prompt_sha"), "extracted_at": ex.get("extracted_at"),
                               "text_truncated": ex.get("truncated") is True, "condensed": ex.get("condensed") is True,
                               "chars_read": ex.get("chars_read", ex.get("n_chars")),
                               "max_chars": ex.get("max_chars", _LEGACY_MAX_CHARS if ex.get("truncated") else None)},
            }, ensure_ascii=False))
    body = ("\n".join(lines) + ("\n" if lines else "")).encode("utf-8")
    return Response(content=body, media_type="application/x-ndjson; charset=utf-8", headers={
        "Content-Disposition": f'attachment; filename="extraction_dataset_{re.sub(r"[^A-Za-z0-9_-]", "_", scenario_id)}.jsonl"',
        "X-Rows": str(len(lines))})


@app.get("/user-scenarios/{scenario_id}/extraction/export")
def export_scenario_extraction(scenario_id: str, format: str = Query("xlsx"),
                               id_prefix: str = Query("LR", max_length=12),
                               include_rejected: bool = Query(False)) -> Response:
    """The extractions of ALL the relevant articles, as the template workbook (`xlsx`) or
    one flat table (`csv`). Articles not extracted yet are not in it; the README sheet and
    the `X-Coverage` header say how many that is."""
    _get_user_scenario_or_404(scenario_id)
    fmt = (format or "xlsx").lower()
    if fmt not in ("xlsx", "csv"):
        raise HTTPException(status_code=400, detail="format must be xlsx or csv")
    rows = _relevant_rows(scenario_id)
    done = [r for r in rows if isinstance(r.get("extraction_json"), dict)]
    from_abs = sum(1 for r in done if r["extraction_json"].get("source") == "abstract")
    n_cut, n_condensed, _dropped = _read_counts(done)
    note = (f"{len(done)} of the {len(rows)} relevant articles are extracted "
            f"({len(rows) - len(done)} not yet); {from_abs} of them from the abstract only; "
            f"{n_cut} read from a text cut at the size limit; {n_condensed} condensed.")
    prefix = re.sub(r"[^A-Za-z0-9_-]", "", id_prefix) or "LR"
    reviews = load_reviews([int(r["id"]) for r in done])
    if fmt == "csv":
        body, media, ext = (build_long_csv(done, prefix, get_index(scenario_id), reviews, include_rejected)
                            .encode("utf-8-sig"), "text/csv; charset=utf-8", "csv")
    else:
        # The workbook also holds the lists the report and the panel cut: every country, every
        # pooled label (those with too few studies too) and every comparison.
        from .pooling import pooled_estimates
        body, media, ext = (build_workbook(done, note, prefix, get_index(scenario_id), reviews, include_rejected,
                                           get_nuts_index(), geo=geography(scenario_id),
                                           pooled=pooled_estimates(scenario_id, full=True)),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx")
    return Response(content=body, media_type=media, headers={
        "Content-Disposition": f'attachment; filename="extraction_{re.sub(r"[^A-Za-z0-9_-]", "_", scenario_id)}.{ext}"',
        "X-Coverage": note,
    })
